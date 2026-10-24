from flask import Blueprint, render_template, request, jsonify, session, send_file
from datetime import datetime
from io import BytesIO
import re
import models
import auth
import utils
from auth import _require_admin_json, _get_current_enrollment_year

bp = Blueprint('lifecycle', __name__)


@bp.route('/admin/student-change', methods=['GET'])
@auth.admin_required
def admin_student_change():
    """学生异动相关接口。"""
    return render_template('admin/student_change.html')


@bp.route('/admin/student-dropout', methods=['GET'])
@auth.admin_required
def admin_student_dropout():
    """学生退学相关接口。"""
    page = int(request.args.get('page', 1))
    date_from = request.args.get('date_from', '').strip() or None
    date_to = request.args.get('date_to', '').strip() or None
    name = request.args.get('name', '').strip() or None
    result = models.get_dropout_students(page=page, per_page=20,
                                         date_from=date_from, date_to=date_to, name=name)
    return render_template('admin/student_dropout.html',
                          students=result['students'],
                          pagination=result)


@bp.route('/admin/student-dropout/export')
@auth.admin_required
def admin_student_dropout_export():
    """学生退学导出相关接口。"""
    date_from = request.args.get('date_from', '').strip() or None
    date_to = request.args.get('date_to', '').strip() or None
    name = request.args.get('name', '').strip() or None
    students = models.get_dropout_students_all(date_from=date_from, date_to=date_to, name=name)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = '退学学生处理'

    headers = ['学生姓名', '身份证号码', '班级', '年级', '入学年',
               '学生电话', '家长电话', '学籍状态', '退学表提交日期', '处理日期', '异动信息', '处理状态']
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    green_fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    yellow_fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')

    for row_idx, s in enumerate(students, 2):
        ws.cell(row=row_idx, column=1, value=s.get('name', ''))

        id_card = s.get('id_card', '') or ''
        ws.cell(row=row_idx, column=2, value=id_card)

        ws.cell(row=row_idx, column=3, value=s.get('admin_class', ''))
        ws.cell(row=row_idx, column=4, value=s.get('grade', '') or '')
        ws.cell(row=row_idx, column=5, value=s.get('enrollment_year', '') or '')

        ws.cell(row=row_idx, column=6, value=s.get('student_phone', '') or '')
        ws.cell(row=row_idx, column=7, value=s.get('parent_phone', '') or '')
        ws.cell(row=row_idx, column=8, value=s.get('registration_status', '') or '')

        dropout_date = s.get('dropout_date', '')
        if dropout_date:
            dropout_date = dropout_date[:10]
        ws.cell(row=row_idx, column=9, value=dropout_date or '')

        ws.cell(row=row_idx, column=11, value=s.get('change_info', '') or '')

        is_deleted = s.get('is_deleted', 0)
        if is_deleted:
            process_date = s.get('process_date', '')
            if process_date:
                process_date = process_date[:10]
            ws.cell(row=row_idx, column=10, value=process_date or '')
            cell = ws.cell(row=row_idx, column=12, value='已处理')
            cell.fill = green_fill
        else:
            ws.cell(row=row_idx, column=10, value='')
            cell = ws.cell(row=row_idx, column=12, value='待处理')
            cell.fill = yellow_fill

    for col in range(1, 13):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

    ws.auto_filter.ref = ws.dimensions

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M')
    filename = f'退学学生处理_{timestamp}.xlsx'

    return send_file(output,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


@bp.route('/admin/graduation')
@auth.admin_required
def admin_graduation():
    """毕业。"""
    page = int(request.args.get('page', 1))
    grade = request.args.get('grade', '').strip() or None
    result = models.get_graduated_students(page=page, per_page=20, grade=grade)
    grades = models.get_graduated_list_grades()
    return render_template('admin/graduation.html',
                          students=result['students'],
                          pagination=result,
                          list_grades=grades,
                          current_list_grade=grade)


@bp.route('/admin/graduation/export')
@auth.admin_required
def admin_graduation_export():
    """毕业导出。"""
    grade = request.args.get('grade', '').strip() or None
    students = models.get_graduated_students_all(grade=grade)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = '毕业学生列表'

    headers = ['学生姓名', '身份证号码', '班级', '年级', '专业', '入学年',
               '学生电话', '家长电话', '学籍状态', '在校状态', '备注']
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF', size=11)

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_idx, s in enumerate(students, 2):
        ws.cell(row=row_idx, column=1, value=s.get('name', ''))
        ws.cell(row=row_idx, column=2, value=s.get('id_card', '') or '')
        ws.cell(row=row_idx, column=3, value=s.get('admin_class', ''))
        ws.cell(row=row_idx, column=4, value=s.get('grade', '') or '')
        ws.cell(row=row_idx, column=5, value=s.get('major', '') or '')
        ws.cell(row=row_idx, column=6, value=s.get('enrollment_year', '') or '')
        ws.cell(row=row_idx, column=7, value=s.get('student_phone', '') or '')
        ws.cell(row=row_idx, column=8, value=s.get('parent_phone', '') or '')
        ws.cell(row=row_idx, column=9, value=s.get('registration_status', '') or '')
        ws.cell(row=row_idx, column=10, value=s.get('attendance_status', '') or '')
        ws.cell(row=row_idx, column=11, value=s.get('remark', '') or '')

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 18

    ws.auto_filter.ref = ws.dimensions

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    timestamp = datetime.now().strftime('%Y%m%d_%H%M')
    filename = f'毕业学生列表_{timestamp}.xlsx'

    return send_file(output,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)


@bp.route('/api/admin/graduation/grades')
@auth.admin_required
def api_graduation_grades():
    """毕业年级。"""
    grades = models.get_graduation_distinct_grades()
    return jsonify({'grades': grades})


@bp.route('/api/admin/graduation/preview/<grade>')
@auth.admin_required
def api_graduation_preview(grade):
    """毕业预览。 参数: grade"""
    admin_class = request.args.get('admin_class', '').strip() or None
    students = models.get_graduation_eligible_students(grade, admin_class)
    return jsonify({'students': students, 'count': len(students)})


@bp.route('/api/admin/graduation/classes/<grade>')
@auth.admin_required
def api_graduation_classes(grade):
    """毕业班级。 参数: grade"""
    classes = models.get_graduation_eligible_classes(grade)
    return jsonify({'classes': classes})


@bp.route('/api/admin/graduation/process', methods=['POST'])
@auth.admin_required
def api_graduation_process():
    """毕业处理。"""
    data = request.get_json()
    if not data:
        return jsonify({'success': False, 'message': '无效数据'}), 400

    student_ids = data.get('student_ids', [])
    if not student_ids:
        return jsonify({'success': False, 'message': '未选择任何学生'}), 400

    if len(student_ids) > 500:
        return jsonify({'success': False, 'message': '单次最多处理500名学生'}), 400

    operator = session.get('display_name', session.get('username', 'unknown'))
    processed, details = models.process_graduation(student_ids, operator)

    return jsonify({
        'success': True,
        'processed': processed,
        'details': details,
        'message': f'成功处理 {processed} 名学生毕业'
    })


@bp.route('/api/admin/graduation/undo', methods=['POST'])
@auth.admin_required
def api_graduation_undo():
    """毕业撤销。"""
    data = request.get_json()
    if not data:
        return jsonify({'success': False, 'message': '无效数据'}), 400

    student_ids = data.get('student_ids', [])
    if not student_ids:
        return jsonify({'success': False, 'message': '未选择任何学生'}), 400

    if len(student_ids) > 500:
        return jsonify({'success': False, 'message': '单次最多处理500名学生'}), 400

    operator = session.get('display_name', session.get('username', 'unknown'))
    try:
        processed, failed, details = models.undo_graduation(student_ids, operator)
        return jsonify({
            'success': True,
            'processed': processed,
            'failed': failed,
            'details': details,
            'message': f'成功撤销 {processed} 人' + (f'，{failed} 人跳过' if failed else '')
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'撤销失败: {str(e)}'}), 500


@bp.route('/admin/graduation/query')
@auth.login_required
def admin_graduation_query():
    """毕业查询。"""
    return render_template('admin/graduation_query.html')


# ========== 实习处理 API ==========
@bp.route('/api/admin/internship/grades')
@auth.admin_required
def api_internship_grades():
    """实习年级。"""
    grades = models.get_internship_grades()
    return jsonify({'grades': grades})


@bp.route('/api/admin/internship/preview/<grade>')
@auth.admin_required
def api_internship_preview(grade):
    """实习预览。 参数: grade"""
    admin_class = request.args.get('admin_class', '').strip()
    students = models.get_internship_eligible_students(grade, admin_class=admin_class if admin_class else None)
    return jsonify({'students': students, 'count': len(students)})


@bp.route('/api/admin/internship/classes/<grade>')
@auth.admin_required
def api_internship_classes(grade):
    """实习班级。 参数: grade"""
    classes = models.get_internship_eligible_classes(grade)
    return jsonify({'classes': classes})


@bp.route('/api/admin/internship/process', methods=['POST'])
@auth.admin_required
def api_internship_process():
    """实习处理。"""
    data = request.get_json()
    if not data:
        return jsonify({'success': False, 'message': '无效数据'}), 400

    student_ids = data.get('student_ids', [])
    if not student_ids:
        return jsonify({'success': False, 'message': '未选择任何学生'}), 400

    if len(student_ids) > 500:
        return jsonify({'success': False, 'message': '单次最多处理500名学生'}), 400

    operator = session.get('display_name', session.get('username', 'unknown'))
    processed, details = models.process_internship(student_ids, operator)

    return jsonify({
        'success': True,
        'processed': processed,
        'details': details,
        'message': f'成功处理 {processed} 名学生的实习状态'
    })


@bp.route('/api/admin/internship/undo', methods=['POST'])
@auth.admin_required
def api_internship_undo():
    """实习撤销。"""
    data = request.get_json()
    if not data:
        return jsonify({'success': False, 'message': '无效数据'}), 400

    student_ids = data.get('student_ids', [])
    if not student_ids:
        return jsonify({'success': False, 'message': '未选择任何学生'}), 400

    if len(student_ids) > 500:
        return jsonify({'success': False, 'message': '单次最多处理500名学生'}), 400

    operator = session.get('display_name', session.get('username', 'unknown'))
    try:
        processed, failed, details = models.undo_internship(student_ids, operator)
        return jsonify({
            'success': True,
            'processed': processed,
            'failed': failed,
            'details': details,
            'message': f'成功撤销 {processed} 人' + (f'，{failed} 人跳过' if failed else '')
        })
    except Exception as e:
        return jsonify({'success': False, 'message': f'撤销失败: {str(e)}'}), 500


@bp.route('/api/admin/internship/processed/grades')
@auth.admin_required
def api_internship_processed_grades():
    """实习已处理年级。"""
    grades = models.get_internship_processed_grades()
    return jsonify({'grades': grades})


@bp.route('/api/admin/internship/processed/<grade>')
@auth.admin_required
def api_internship_processed(grade):
    """实习已处理。 参数: grade"""
    students = models.get_internship_processed_students(grade)
    return jsonify({'students': students, 'count': len(students)})


@bp.route('/api/admin/graduation/filters')
@auth.login_required
def api_graduation_filters():
    """返回毕业生可用的筛选项（学校类别/年级/班级），支持按已选条件联动"""
    school = request.args.get('school', '').strip()
    grade = request.args.get('grade', '').strip()

    # 数据权限校验：显式传入的学校/年级必须在本用户可见范围内
    if school and not auth.check_school_permission(school):
        return jsonify({'grades': [], 'classes': []})
    if grade and not auth.check_grade_permission(grade):
        return jsonify({'grades': [], 'classes': []})

    conditions = [
        "s.registration_status = '毕业'",
        "s.is_deleted = 0"
    ]
    params = []

    if school:
        conditions.append('s.school = ?')
        params.append(int(school))
    if grade:
        conditions.append('s.grade = ?')
        params.append(grade)

    where = ' AND '.join(conditions)

    conn = models.get_db()
    result = {'grades': [], 'classes': []}

    if not grade:
        grades = conn.execute(
            f'''SELECT DISTINCT s.grade FROM students s
                WHERE {where} AND s.grade IS NOT NULL AND s.grade != ''
                ORDER BY s.grade DESC''',
            params
        ).fetchall()
        result['grades'] = [r[0] for r in grades]

    if grade:
        classes = conn.execute(
            f'''SELECT DISTINCT s.admin_class FROM students s
                WHERE {where} AND s.admin_class IS NOT NULL AND s.admin_class != ''
                ORDER BY s.admin_class''',
            params
        ).fetchall()
        result['classes'] = [r[0] for r in classes]

    conn.close()
    return jsonify(result)


@bp.route('/api/admin/graduation/query')
@auth.login_required
def api_graduation_query():
    # 应用权限筛选（学校可见性、年级可见性）
    """毕业查询。"""
    permission_filters = auth.apply_permission_filters({})

    user_school = request.args.get('school', '').strip()
    school = user_school or permission_filters.get('school')
    user_grade = request.args.get('grade', '').strip()
    grade = user_grade or permission_filters.get('grade')
    admin_class = request.args.get('admin_class', '').strip()
    name = request.args.get('name', '').strip()

    # 学校/年级数据权限校验：仅校验用户显式传入的值（权限注入值已在可见范围内）
    if user_school and not auth.check_school_permission(user_school):
        return jsonify({'students': [], 'count': 0, 'visible_fields': auth.get_visible_fields()})
    if user_grade and not auth.check_grade_permission(user_grade):
        return jsonify({'students': [], 'count': 0, 'visible_fields': auth.get_visible_fields()})

    conditions = [
        "s.registration_status = '毕业'",
        "s.is_deleted = 0"
    ]
    params = []

    if school:
        school_val = int(school) if isinstance(school, str) else school
        if isinstance(school_val, list):
            placeholders = ','.join(['?' for _ in school_val])
            conditions.append(f's.school IN ({placeholders})')
            params.extend(school_val)
        else:
            conditions.append('s.school = ?')
            params.append(school_val)
    if grade:
        if isinstance(grade, list):
            placeholders = ','.join(['?' for _ in grade])
            conditions.append(f's.grade IN ({placeholders})')
            params.extend(grade)
        else:
            conditions.append('s.grade = ?')
            params.append(grade)
    if admin_class:
        conditions.append('s.admin_class = ?')
        params.append(admin_class)
    if name:
        conditions.append('(s.name LIKE ? OR s.name_pinyin LIKE ?)')
        params.append(f'%{name}%')
        params.append(f'%{name}%')

    where = ' AND '.join(conditions)

    # 档案派遣权限分两档：
    # 1) 派遣状态：管理员默认可见；普通用户需在用户管理「可见字段」中勾选「派遣状态」
    # 2) 派遣明细（批次/机构/电话/日期等）：与档案派遣模块一致，仅管理员可见
    is_admin = session.get('role') == 'admin'
    can_view_dispatch = is_admin or auth.check_field_permission('派遣状态')
    can_view_dispatch_detail = is_admin
    dispatch_join = ''
    dispatch_fields = ''
    if can_view_dispatch:
        dispatch_join = ('LEFT JOIN archive_dispatch d '
                         'ON d.student_id = s.id AND d.is_deleted = 0')
        dispatch_fields = ''',
                   d.id AS dispatch_id, d.dispatch_status'''
        if can_view_dispatch_detail:
            dispatch_fields += ''',
                   d.dispatch_no, d.batch_no, d.dispatch_date, d.returned_at,
                   d.agency_name, d.agency_address, d.agency_phone,
                   d.match_status AS dispatch_match_status,
                   d.operator AS dispatch_operator,
                   d.remark AS dispatch_remark'''

    conn = models.get_db()
    rows = conn.execute(
        f'''SELECT s.id, s.name, s.student_no, s.national_no, s.id_card, s.grade,
                   s.admin_class, s.major, s.attendance_status,
                   s.registration_status, s.remark, s.enrollment_year,
                   s.gender, s.student_phone, s.parent_phone, s.school
                   {dispatch_fields}
            FROM students s
            {dispatch_join}
            WHERE {where}
            ORDER BY s.grade DESC, s.admin_class, s.name
            LIMIT 500''',
        params
    ).fetchall()
    conn.close()

    # 字段级权限过滤
    raw_students = [dict(r) for r in rows]
    students = [auth.filter_student_data(s) for s in raw_students]
    # 派遣字段不属于学生字段映射，字段权限过滤会将其剔除，这里按已授权范围补回
    if can_view_dispatch:
        dispatch_keys = ['dispatch_id', 'dispatch_status']
        if can_view_dispatch_detail:
            dispatch_keys += ['dispatch_no', 'batch_no', 'dispatch_date', 'returned_at',
                              'agency_name', 'agency_address', 'agency_phone',
                              'dispatch_match_status', 'dispatch_operator', 'dispatch_remark']
        for src, dst in zip(raw_students, students):
            for key in dispatch_keys:
                if key in src:
                    dst[key] = src[key]
    visible_fields = auth.get_visible_fields()

    # 构建列配置：映射中文名 → 显示信息，前端据此动态渲染
    COLUMN_LABELS = {
        '学生姓名': 'name', '学号': 'student_no', '学生身份证号': 'id_card',
        '行政班级': 'admin_class', '年级': 'grade', '专业': 'major',
        '学校类别': 'school', '学籍状态': 'registration_status',
        '在校状态': 'attendance_status', '性别': 'gender',
        '学生电话': 'phone', '家长电话': 'phone', '备注': 'remark',
        '国网学籍号': 'national_no',
    }
    # 界面显示名：权限字段名为 '国网学籍号'，页面显示为 '学籍号'
    COLUMN_LABELS_DISPLAY = {'国网学籍号': '学籍号'}
    COLUMN_ORDER = ['学生姓名', '学号', '国网学籍号', '学生身份证号', '行政班级', '年级',
                    '专业', '学校类别', '学籍状态', '在校状态', '性别',
                    '学生电话', '家长电话', '备注']
    columns = []
    seen_keys = set()
    for cn in COLUMN_ORDER:
        if cn in visible_fields or visible_fields == ['*']:
            key = COLUMN_LABELS.get(cn)
            if key and key not in seen_keys:
                seen_keys.add(key)
                columns.append({'key': key, 'label': COLUMN_LABELS_DISPLAY.get(cn, cn)})

    # 派遣状态列（管理员或被授权「派遣状态」字段的用户）：追加在表格末尾；
    # 导出时管理员额外带上批次/机构/日期等明细列
    export_columns = list(columns)
    if can_view_dispatch:
        columns.append({'key': 'dispatch_status', 'label': '派遣状态'})
        export_columns.append({'key': 'dispatch_status', 'label': '派遣状态'})
    if can_view_dispatch_detail:
        export_columns.extend([
            {'key': 'batch_no', 'label': '派遣批次'},
            {'key': 'agency_name', 'label': '接收机构'},
            {'key': 'agency_phone', 'label': '机构电话'},
            {'key': 'dispatch_date', 'label': '派遣日期'},
            {'key': 'returned_at', 'label': '退档日期'},
        ])

    return jsonify({
        'students': students,
        'count': len(students),
        'visible_fields': visible_fields,
        'columns': columns,
        'export_columns': export_columns,
        'dispatch_visible': can_view_dispatch,
        'dispatch_detail_visible': can_view_dispatch_detail,
    })


@bp.route('/api/students/search-quick', methods=['GET'])
def api_student_search_quick():
    """学生搜索quick相关接口。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    keyword = request.args.get('keyword', '').strip()
    by_id_card = request.args.get('by_id_card') == '1'
    if not keyword:
        return jsonify({'results': [], 'total': 0})

    conn = models.get_db()
    cursor = conn.cursor()

    if by_id_card:
        cursor.execute('''
            SELECT id, student_no, id_card, name, gender, major, admin_class,
                   attendance_status, registration_status, change_info
            FROM students WHERE id_card = ? AND is_deleted = 0
        ''', (keyword,))
    elif keyword.isdigit() and len(keyword) == 4:
        cursor.execute('''
            SELECT id, student_no, id_card, name, gender, major, admin_class,
                   attendance_status, registration_status, change_info
            FROM students WHERE id_card LIKE ? AND is_deleted = 0
            ORDER BY student_no
            LIMIT 50
        ''', ('%' + keyword,))
    else:
        cursor.execute('''
            SELECT id, student_no, id_card, name, gender, major, admin_class,
                   attendance_status, registration_status, change_info
            FROM students WHERE (name LIKE ? OR name_pinyin LIKE ?) AND is_deleted = 0
            ORDER BY student_no
            LIMIT 50
        ''', ('%' + keyword + '%', '%' + keyword + '%'))

    rows = cursor.fetchall()
    conn.close()

    results = []
    for row in rows:
        results.append({
            'id': row['id'],
            'student_no': row['student_no'] or '',
            'id_card': row['id_card'] or '',
            'name': row['name'] or '',
            'gender': row['gender'] or '',
            'major': row['major'] or '',
            'admin_class': row['admin_class'] or '',
            'attendance_status': row['attendance_status'] or '',
            'registration_status': row['registration_status'] or '',
            'change_info': row['change_info'] or ''
        })

    return jsonify({'results': results, 'total': len(results)})


@bp.route('/api/students/class-options', methods=['GET'])
def api_student_class_options():
    """学生班级options相关接口。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    conn = models.get_db()
    cursor = conn.cursor()
    cursor.execute('''
        SELECT DISTINCT admin_class FROM students
        WHERE is_deleted = 0 AND admin_class IS NOT NULL AND admin_class != ''
        ORDER BY admin_class
    ''')
    rows = cursor.fetchall()
    conn.close()
    classes = [row['admin_class'] for row in rows]
    return jsonify({'classes': classes})


@bp.route('/api/students/change-record', methods=['POST'])
def api_student_change_record():
    """学生异动记录相关接口。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    data = request.get_json()
    action = data.get('action', '').strip()
    id_card = data.get('id_card', '').strip()
    new_class = data.get('new_class', '').strip()
    new_major = data.get('new_major', '').strip()
    change_info = data.get('change_info', '').strip()

    if not id_card or not action:
        return jsonify({'success': False, 'message': '参数不完整'}), 400

    student = models.get_student_by_id_card(id_card)
    if not student:
        return jsonify({'success': False, 'message': '未找到该学生'}), 404

    operator = session.get('display_name', session.get('username', 'unknown'))

    try:
        if action == 'class_change':
            if not new_class or not change_info:
                return jsonify({'success': False, 'message': '参数不完整'}), 400
            old_class = student.get('admin_class', '') or ''
            old_change_info = student.get('change_info', '') or ''
            models.update_student(student['id'], {
                'admin_class': new_class,
                'change_info': change_info
            })
            models.log_update(student['id'], id_card, 'admin_class', old_class, new_class, operator)
            models.log_update(student['id'], id_card, 'change_info', old_change_info, change_info, operator)
            models.update_sync_timestamp()
            return jsonify({
                'success': True,
                'message': student['name'] + ' 班级已由 ' + old_class + ' 变更为 ' + new_class,
                'student': {
                    'admin_class': new_class,
                    'major': student.get('major', ''),
                    'attendance_status': student.get('attendance_status', ''),
                    'change_info': change_info
                }
            })

        elif action == 'major_change':
            if not new_class or not new_major or not change_info:
                return jsonify({'success': False, 'message': '参数不完整'}), 400
            old_class = student.get('admin_class', '') or ''
            old_major = student.get('major', '') or ''
            old_change_info = student.get('change_info', '') or ''

            update_data = {'change_info': change_info}
            update_data['admin_class'] = new_class
            update_data['major'] = new_major
            if new_major != old_major:
                frontend_school = data.get('school')
                if frontend_school and int(frontend_school) > 0:
                    update_data['school'] = int(frontend_school)
                else:
                    update_data['school'] = utils._derive_school(new_major)[0]
            models.update_student(student['id'], update_data)

            if new_class != old_class:
                models.log_update(student['id'], id_card, 'admin_class', old_class, new_class, operator)
            if new_major != old_major:
                models.log_update(student['id'], id_card, 'major', old_major, new_major, operator)
            models.log_update(student['id'], id_card, 'change_info', old_change_info, change_info, operator)
            models.update_sync_timestamp()

            return jsonify({
                'success': True,
                'message': student['name'] + ' 专业/班级变更成功',
                'student': {
                    'admin_class': new_class,
                    'major': new_major,
                    'attendance_status': student.get('attendance_status', ''),
                    'change_info': change_info
                }
            })

        elif action == 'dropout':
            if not change_info:
                return jsonify({'success': False, 'message': '参数不完整'}), 400
            old_status = student.get('attendance_status', '') or ''
            old_change_info = student.get('change_info', '') or ''
            models.update_student(student['id'], {
                'attendance_status': '退学办结',
                'change_info': change_info
            })
            models.log_update(student['id'], id_card, 'attendance_status', old_status, '退学办结', operator)
            models.log_update(student['id'], id_card, 'change_info', old_change_info, change_info, operator)
            models.update_sync_timestamp()
            return jsonify({
                'success': True,
                'message': student['name'] + ' 已办理退学办结',
                'student': {
                    'admin_class': student.get('admin_class', ''),
                    'major': student.get('major', ''),
                    'attendance_status': '退学办结',
                    'change_info': change_info
                }
            })

        else:
            return jsonify({'success': False, 'message': '未知操作类型'}), 400

    except Exception as e:
        return jsonify({'success': False, 'message': '操作失败：' + str(e)}), 500


@bp.route('/api/transfer-pending/students')
def api_transfer_pending_students():
    """转学待处理学生。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    page = int(request.args.get('page', 1))
    result = models.get_transfer_pending_students(page=page, per_page=10)
    return jsonify(result)


@bp.route('/api/students/check-student-no')
def api_check_student_no():
    """检查学生编号。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    student_no = request.args.get('student_no', '').strip()
    exclude_id = request.args.get('exclude_id', '').strip()
    if not student_no:
        return jsonify({'exists': False, 'message': '学号不能为空'})

    existing = models.get_student_by_student_no(student_no)
    if existing:
        if exclude_id and str(existing.get('id', '')) == exclude_id:
            return jsonify({'exists': False, 'message': ''})
        return jsonify({'exists': True, 'message': '学号已存在：' + (existing.get('name', '') or '')})
    return jsonify({'exists': False, 'message': ''})


@bp.route('/api/transfer-pending/process', methods=['POST'])
def api_transfer_pending_process():
    """转学待处理处理。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    data = request.get_json()
    student_id = data.get('student_id')
    student_no = data.get('student_no', '').strip()

    if not student_id or not student_no:
        return jsonify({'success': False, 'message': '参数不完整'}), 400

    if not re.match(r'^\d{6}$', student_no):
        return jsonify({'success': False, 'message': '学号需为6位数字'}), 400

    existing = models.get_student_by_student_no(student_no)
    if existing and str(existing.get('id', '')) != str(student_id):
        return jsonify({'success': False, 'message': '学号已被占用'}), 409

    conn = models.get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT registration_status, name, id_card FROM students WHERE id = ? AND is_deleted = 0', (student_id,))
    student = cursor.fetchone()
    if not student:
        conn.close()
        return jsonify({'success': False, 'message': '学生不存在'}), 404

    old_registration = student['registration_status'] or ''
    if old_registration != '未注册':
        conn.close()
        return jsonify({'success': False, 'message': '该学生学籍状态不是未注册，当前状态：' + old_registration}), 400

    operator = session.get('display_name', session.get('username', 'unknown'))

    cursor.execute('''
        UPDATE students SET student_no = ?, registration_status = '正式学籍'
        WHERE id = ? AND is_deleted = 0
    ''', (student_no, student_id))
    conn.commit()
    conn.close()

    models.log_update(student_id, student['id_card'], 'student_no', '', student_no, operator)
    models.log_update(student_id, student['id_card'], 'registration_status', old_registration, '正式学籍', operator)
    models.update_sync_timestamp()

    return jsonify({
        'success': True,
        'message': student['name'] + ' 转学处理完成，学籍状态已变更为正式学籍'
    })
