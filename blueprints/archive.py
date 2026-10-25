# -*- coding: utf-8 -*-
"""档案派遣：人才机构字典、按县批次生成、转递单导出、升学标记与转递单配置。"""
import os
from datetime import datetime

from flask import Blueprint, render_template, request, jsonify, session, send_file
from io import BytesIO

import auth
import config
import models

bp = Blueprint('archive', __name__)

TRANSFER_TEMPLATE_PATH = os.path.join(config.BASE_DIR, 'archive_templates', '档案转递单模板.docx')

# 模板占位符 → 数据字段
TRANSFER_MAPPINGS = [
    ('student_no', 'student_no'), ('graduate_year', 'graduate_year'),
    ('student_name', 'student_name'), ('id_card', 'id_card'),
    ('origin_place', 'origin_place'), ('student_phone', 'student_phone'),
    ('parent_phone', 'parent_phone'), ('agency_name', 'agency_name'),
    ('agency_address', 'agency_address'), ('agency_phone', 'agency_phone'),
    ('dispatch_year', 'dispatch_year'), ('dispatch_month', 'dispatch_month'),
    ('dispatch_day', 'dispatch_day'), ('handler', 'handler'),
    ('handler_phone', 'handler_phone'), ('school_name', 'school_name'),
    ('receipt_address', 'receipt_address'), ('receipt_postal', 'receipt_postal'),
    ('archive_volume', 'archive_volume'), ('employer_name', 'employer_name'),
    ('employer_address', 'employer_address'), ('dispatch_no', 'dispatch_no'),
]


def _operator():
    return session.get('display_name', session.get('username', 'unknown'))


def _mappings():
    return [{'placeholder': '${' + key + '}', 'source_field': field, 'default_value': ''}
            for key, field in TRANSFER_MAPPINGS]


# ---------------------------------------------------------------- 页面
@bp.route('/admin/archive/agencies')
@auth.admin_required
def agencies_page():
    return render_template('archive/agencies.html')


@bp.route('/admin/archive/dispatch')
@auth.admin_required
def dispatch_page():
    grades = models.get_graduated_list_grades()
    return render_template('archive/dispatch.html', list_grades=grades)


@bp.route('/admin/archive/settings', methods=['GET'])
@auth.admin_required
def settings_page():
    return render_template('archive/settings.html', settings=models.get_all_settings())


# ---------------------------------------------------------------- 机构字典
@bp.route('/api/archive/agencies')
@auth.admin_required
def api_agencies():
    keyword = request.args.get('keyword', '').strip()
    level = request.args.get('level', '').strip()
    status = request.args.get('status', '').strip()
    rows = models.list_agencies(keyword=keyword, level=level, status=status)
    return jsonify({'agencies': rows, 'count': len(rows),
                    'coverage': models.coverage(request.args.get('grade') or None)})


@bp.route('/api/archive/agencies', methods=['POST'])
@auth.admin_required
def api_agency_create():
    data = request.get_json(silent=True) or {}
    try:
        agency = models.create_agency(data)
    except ValueError as exc:
        return jsonify({'success': False, 'message': str(exc)}), 400
    return jsonify({'success': True, 'agency': agency,
                    'message': f'已新增机构「{agency["agency_name"]}」'})


@bp.route('/api/archive/agencies/import', methods=['POST'])
@auth.admin_required
def api_agencies_import():
    upload = request.files.get('file')
    if not upload or not upload.filename:
        return jsonify({'success': False, 'message': '请选择 Excel 文件'}), 400
    if not upload.filename.lower().endswith(('.xlsx', '.xlsm')):
        return jsonify({'success': False, 'message': '仅支持 .xlsx / .xlsm'}), 400
    try:
        stats = models.import_agencies(upload)
    except Exception as exc:  # noqa: BLE001
        return jsonify({'success': False, 'message': f'导入失败：{exc}'}), 400
    return jsonify({
        'success': True,
        'stats': stats,
        'message': (f"新增 {stats['inserted']} 条、更新 {stats['updated']} 条、"
                    f"特殊区域 {stats['special']} 条、冲突 {stats['conflict']} 条")
    })


@bp.route('/api/archive/agencies/<int:agency_id>/toggle', methods=['POST'])
@auth.admin_required
def api_agency_toggle(agency_id):
    new_status = models.toggle_agency(agency_id)
    if new_status is None:
        return jsonify({'success': False, 'message': '机构不存在'}), 404
    return jsonify({'success': True, 'status': new_status})


@bp.route('/admin/archive/agencies/template')
@auth.admin_required
def agencies_template():
    from openpyxl import Workbook

    headers, sample = models.agency_template_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = '人才机构'
    ws.append(headers)
    for row in sample:
        ws.append(row)
    for idx in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = 22
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True,
                     download_name='人才机构导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ---------------------------------------------------------------- 派遣工作台
@bp.route('/api/archive/overview')
@auth.admin_required
def api_overview():
    grade = request.args.get('grade', '').strip() or None
    school = request.args.get('school', '').strip() or None
    rows = models.load_graduates(grade, school)
    counties = models.summarize_counties(rows)
    return jsonify({
        'counties': counties,
        'batches': models.list_batches(school),
        'settings': models.get_all_settings(),
        'coverage': models.coverage(grade),
    })


@bp.route('/api/archive/students')
@auth.admin_required
def api_students():
    grade = request.args.get('grade', '').strip() or None
    county = request.args.get('county', '').strip()
    status = request.args.get('status', '').strip()
    school = request.args.get('school', '').strip() or None
    rows = models.load_graduates(grade, school)

    if county:
        rows = [r for r in rows if r['group_code'] == county]

    def keep(row):
        current = row['dispatch_status']
        matched = bool(row['matched_agency'])
        if status == 'pending':
            return current in ('', models.STATUS_PENDING, models.STATUS_RETURNED) and matched
        if status == 'unmatched':
            return current in ('', models.STATUS_PENDING, models.STATUS_RETURNED) and not matched
        if status == 'sent':
            return current == models.STATUS_SENT
        if status == 'returned':
            return current == models.STATUS_RETURNED
        if status == 'study':
            return current == models.STATUS_STUDY
        return True

    rows = [r for r in rows if keep(r)]
    data = []
    for row in rows:
        data.append({
            'id': row['id'],
            'student_no': row.get('student_no') or '',
            'name': row.get('name') or '',
            'id_card': row.get('id_card') or '',
            'school': row.get('school'),
            'grade': row.get('grade') or '',
            'admin_class': row.get('admin_class') or '',
            'county_code': row['group_code'],
            'county_name': row['group_name'],
            'graduate_year': models.graduate_year_of(row.get('grade')),
            'dispatch_id': row.get('dispatch_id'),
            'dispatch_status': row['dispatch_status'],
            'batch_no': row.get('batch_no') or '',
            'match_status': row['match_status'],
            'match_reason': row['match_reason'],
            'agency_name': (row.get('matched_agency') or {}).get('agency_name', '') or row.get('agency_name') or '',
            'agency_address': (row.get('matched_agency') or {}).get('address', '') or row.get('agency_address') or '',
            'agency_phone': (row.get('matched_agency') or {}).get('phone', '') or row.get('agency_phone') or '',
            'student_phone': row.get('student_phone') or '',
            'parent_phone': row.get('parent_phone') or '',
            'origin_place': row.get('origin_place') or '',
        })
    return jsonify({'students': data, 'count': len(data)})


@bp.route('/api/archive/generate', methods=['POST'])
@auth.admin_required
def api_generate():
    data = request.get_json() or {}
    county_code = data.get('county_code', '').strip()
    student_ids = data.get('student_ids') or []
    if not county_code:
        return jsonify({'success': False, 'message': '请选择县'}), 400
    if not student_ids:
        return jsonify({'success': False, 'message': '请勾选学生'}), 400
    if len(student_ids) > 500:
        return jsonify({'success': False, 'message': '单批最多 500 人'}), 400
    try:
        result = models.create_batch(county_code, student_ids, _operator(),
                                     data.get('overrides') or {})
    except ValueError as exc:
        return jsonify({'success': False, 'message': str(exc)}), 400
    return jsonify({'success': True, **result,
                    'message': f"批次 {result['batch_no']} 已生成，共 {result['inserted']} 人"})


@bp.route('/admin/archive/export/<path:batch_no>')
@auth.admin_required
def export_batch(batch_no):
    from blueprints.export import _export_docx

    # ?redispatch=1：只打印「曾退档且已重新派遣」的重派单，避免同批其他人被重复打印
    redispatch_only = request.args.get('redispatch') == '1'
    rows = models.build_export_rows(batch_no, redispatch_only=redispatch_only)
    if not rows:
        message = ('该批次没有可打印的重派转递单'
                   if redispatch_only else
                   '该批次没有可打印的记录（退档待重派 / 已失效 / 升学的档案不出转递单）')
        return jsonify({'success': False, 'message': message}), 404
    if not os.path.exists(TRANSFER_TEMPLATE_PATH):
        return jsonify({'success': False, 'message': '转递单模板缺失，请检查 archive_templates 目录'}), 500

    path = _export_docx(TRANSFER_TEMPLATE_PATH, rows, _mappings())
    title = '档案转递单（重派）' if redispatch_only else '档案转递单'
    models.add_export_history(None, title, batch_no, len(rows), path, _operator())
    suffix = '_重派转递单' if redispatch_only else '_档案转递单'
    return send_file(path, as_attachment=True, download_name=f'{batch_no}{suffix}.docx')


@bp.route('/api/archive/dispatch/<int:dispatch_id>/status', methods=['POST'])
@auth.admin_required
def api_dispatch_status(dispatch_id):
    data = request.get_json() or {}
    try:
        models.update_dispatch(dispatch_id, data.get('action', ''), data, _operator())
    except ValueError as exc:
        return jsonify({'success': False, 'message': str(exc)}), 400
    return jsonify({'success': True, 'message': '状态已更新'})


@bp.route('/api/archive/batch/<path:batch_no>/status', methods=['POST'])
@auth.admin_required
def api_batch_status(batch_no):
    data = request.get_json() or {}
    try:
        result = models.update_batch_status(batch_no, data.get('action', ''), _operator())
    except ValueError as exc:
        return jsonify({'success': False, 'message': str(exc)}), 400
    message = f"已标记 {result['count']} 人为已派遣"
    if result['skipped_returned']:
        message += f"，跳过 {result['skipped_returned']} 条退档记录（需先「恢复待派遣」）"
    return jsonify({'success': True, **result, 'message': message})


@bp.route('/api/archive/dispatch/<int:dispatch_id>/agency', methods=['POST'])
@auth.admin_required
def api_dispatch_agency(dispatch_id):
    data = request.get_json() or {}
    try:
        models.update_dispatch_agency(dispatch_id, data.get('agency_id'),
                                     data.get('manual') or {})
    except ValueError as exc:
        return jsonify({'success': False, 'message': str(exc)}), 400
    return jsonify({'success': True, 'message': '接收机构已更新'})


@bp.route('/api/archive/agency-search')
@auth.admin_required
def api_agency_search():
    keyword = request.args.get('keyword', '').strip()
    return jsonify({'agencies': models.list_students_for_agency(keyword)})


# ---------------------------------------------------------------- 升学
@bp.route('/api/archive/further-study/mark', methods=['POST'])
@auth.admin_required
def api_mark_study():
    data = request.get_json() or {}
    ids = data.get('student_ids') or []
    if not ids:
        return jsonify({'success': False, 'message': '请勾选学生'}), 400
    result = models.mark_further_study(ids, _operator())
    return jsonify({'success': True, **result,
                    'message': f"已标记 {result['marked']} 人为升学"})


@bp.route('/api/archive/further-study/cancel', methods=['POST'])
@auth.admin_required
def api_cancel_study():
    data = request.get_json() or {}
    ids = data.get('student_ids') or []
    if not ids:
        return jsonify({'success': False, 'message': '请勾选学生'}), 400
    count = models.cancel_further_study(ids)
    return jsonify({'success': True, 'count': count,
                    'message': f'已取消 {count} 人的升学标记'})


@bp.route('/api/archive/further-study')
@auth.admin_required
def api_study_list():
    grade = request.args.get('grade', '').strip() or None
    county = request.args.get('county', '').strip()
    school = request.args.get('school', '').strip() or None
    rows = models.list_further_study(grade=grade, county=county, school=school)
    return jsonify({'students': rows, 'count': len(rows)})


@bp.route('/api/archive/further-study/import', methods=['POST'])
@auth.admin_required
def api_study_import():
    upload = request.files.get('file')
    if not upload or not upload.filename:
        return jsonify({'success': False, 'message': '请选择 Excel 文件'}), 400
    try:
        from openpyxl import load_workbook

        wb = load_workbook(upload, data_only=True, read_only=True)
        ws = wb.active
        raw = list(ws.iter_rows(values_only=True))
        wb.close()
        header_idx, headers = None, []
        for i, row in enumerate(raw[:10]):
            vals = [str(c).strip() if c is not None else '' for c in row]
            if '学号' in vals or '身份证号' in vals:
                header_idx, headers = i, vals
                break
        if header_idx is None:
            return jsonify({'success': False, 'message': '未找到「学号」或「身份证号」表头'}), 400

        def col(name):
            return headers.index(name) if name in headers else None

        c_no, c_id = col('学号'), col('身份证号')
        conn = models.get_db()
        hit_ids, missing = [], []
        for row in raw[header_idx + 1:]:
            values = [str(c).strip() if c is not None else '' for c in row]
            no = values[c_no] if c_no is not None and c_no < len(values) else ''
            id_card = values[c_id] if c_id is not None and c_id < len(values) else ''
            if not no and not id_card:
                continue
            student = None
            if no:
                student = conn.execute(
                    "SELECT id FROM students WHERE student_no = ? AND is_deleted = 0 AND registration_status = '毕业'",
                    (no,)).fetchone()
            if not student and id_card:
                student = conn.execute(
                    "SELECT id FROM students WHERE id_card = ? AND is_deleted = 0 AND registration_status = '毕业'",
                    (id_card,)).fetchone()
            if student:
                hit_ids.append(student['id'])
            else:
                missing.append(no or id_card)
        conn.close()
        result = models.mark_further_study(hit_ids, _operator())
        return jsonify({'success': True, 'marked': result['marked'], 'missing': missing,
                        'skipped': result['skipped'],
                        'message': f"命中 {len(hit_ids)} 人，标记 {result['marked']} 人，未找到 {len(missing)} 人"})
    except Exception as exc:  # noqa: BLE001
        return jsonify({'success': False, 'message': f'导入失败：{exc}'}), 400


@bp.route('/admin/archive/further-study/template')
@auth.admin_required
def download_study_template():
    """下载升学名单导入模板（表头 + 填写说明）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = '升学名单'
    headers = ['学号', '姓名', '身份证号', '录取院校']
    fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    for idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=idx, value=h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 22
    ws.freeze_panes = 'A2'

    guide = wb.create_sheet('填写说明')
    lines = [
        '升学名单导入说明',
        '',
        '1. 文件格式：Excel（.xlsx / .xlsm），读取第一个工作表。',
        '2. 表头要求：前 10 行内需出现「学号」或「身份证号」列，二者至少填一个；',
        '   姓名、录取院校等其它列仅作备注，系统不读取。',
        '3. 匹配范围：仅匹配学籍状态为「毕业」且未删除的学生；学号优先，学号未命中再用身份证号。',
        '4. 导入结果：命中学生标记为「升学」，不再参与档案派遣；',
        '   已是升学状态的学生自动跳过；未匹配到的学号/身份证号会在导入结果中提示。',
        '5. 导入不受页面当前届别/校别筛选影响，请确保名单本身已核对准确。',
        '6. 取消升学：在「升学名单」列表中勾选学生后点击「取消升学」。',
    ]
    for r, text in enumerate(lines, 1):
        cell = guide.cell(row=r, column=1, value=text)
        if r == 1:
            cell.font = Font(bold=True, size=13)
    guide.column_dimensions['A'].width = 100

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='升学名单导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.route('/admin/archive/further-study/export')
@auth.admin_required
def export_study_list():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    grade = request.args.get('grade', '').strip() or None
    county = request.args.get('county', '').strip()
    school = request.args.get('school', '').strip() or None
    rows = models.list_further_study(grade=grade, county=county, school=school)

    wb = Workbook()
    ws = wb.active
    ws.title = '升学名单'
    headers = ['学号', '姓名', '身份证号', '届', '年级', '班级', '县', '专业方向', '标记时间']
    fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    for idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=idx, value=h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[cell.column_letter].width = 18

    for r, item in enumerate(rows, 2):
        ws.cell(row=r, column=1, value=item.get('student_no', ''))
        ws.cell(row=r, column=2, value=item.get('student_name', ''))
        ws.cell(row=r, column=3, value=item.get('id_card', ''))
        ws.cell(row=r, column=4, value=item.get('graduate_year', ''))
        ws.cell(row=r, column=5, value=item.get('grade', ''))
        ws.cell(row=r, column=6, value=item.get('admin_class', ''))
        ws.cell(row=r, column=7, value=item.get('county_name', ''))
        ws.cell(row=r, column=8, value=item.get('major_direction', '') or '')
        ws.cell(row=r, column=9, value=(item.get('updated_at') or '')[:19])

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"升学名单_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(output, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ---------------------------------------------------------------- 转递单配置
@bp.route('/api/archive/settings', methods=['POST'])
@auth.admin_required
def api_save_settings():
    data = request.get_json() or {}
    school = data.get('school')
    if school not in ('1', '2', 1, 2):
        return jsonify({'success': False, 'message': '校别无效'}), 400
    models.save_settings(school, data, _operator())
    return jsonify({'success': True, 'message': '转递单配置已保存'})
