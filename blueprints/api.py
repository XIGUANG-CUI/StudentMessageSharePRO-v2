import os
import re
import base64
import zipfile
from io import BytesIO
from datetime import datetime

from flask import Blueprint, request, session, jsonify, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

import models
import auth
import utils
import config
import pinyin_utils
from auth import _require_admin_json, _require_login_json, _get_current_enrollment_year
from services.card_reader import get_id_card_reader
from utils import _get_card_reader_ws_url, _set_card_reader_ws_url

bp = Blueprint('api', __name__)


@bp.route('/api/school/derive')
def api_derive_school():
    """推导学校。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    major = request.args.get('major', '').strip()
    school, derived = utils._derive_school(major)
    return jsonify({'school': school, 'major': major, 'derived': derived})


# _require_admin_json / _require_login_json 已移至 auth.py（P1-1 公共化）

@bp.route('/api/students')
def api_students():
    """学生。"""
    auth_error = _require_login_json()
    if auth_error:
        return auth_error
    page = int(request.args.get('page', 1))
    per_page = int(request.args.get('per_page', 20))
    
    filters = utils.parse_filters_from_request(request)
    filters = auth.apply_permission_filters(filters)
    
    sort_field = request.args.get('sort_field', 'id')
    sort_order = request.args.get('sort_order', 'DESC')
    
    result = models.get_students(filters, page=page, per_page=per_page, 
                                 sort_field=sort_field, sort_order=sort_order)
    
    visible_fields = auth.get_visible_fields()
    
    students_data = []
    for student in result['students']:
        bd = student.get('birth_date')
        if not bd and student.get('id_card'):
            bd = utils.extract_birth_date_from_id_card(student['id_card'])
        student['age'] = utils.calculate_age(bd) if bd else '-'
        filtered_student = auth.filter_student_data(student)
        
        for key in filtered_student:
            if key in ['id_card', 'parent_id_card', 'parent_phone', 'student_phone']:
                filtered_student[key] = auth.mask_sensitive_data(filtered_student[key], key)
        
        students_data.append(filtered_student)
    
    return jsonify({
        'students': students_data,
        'total': result['total'],
        'page': result['page'],
        'pages': result['pages']
    })

@bp.route('/api/student/<int:student_id>')
def api_student(student_id):
    """学生相关接口。 参数: student_id"""
    auth_error = _require_login_json()
    if auth_error:
        return auth_error
    student = models.get_student_by_id(student_id)
    if not student:
        return jsonify({'error': '学生不存在'}), 404
    
    if not auth.check_school_permission(student['school']):
        return jsonify({'error': '无权限访问'}), 403
    
    student['age'] = utils.calculate_age(student.get('birth_date'))
    filtered_student = auth.filter_student_data(student)
    
    for key in filtered_student:
        if key in ['id_card', 'parent_id_card', 'parent_phone', 'student_phone']:
            filtered_student[key] = auth.mask_sensitive_data(filtered_student[key], key)
    
    return jsonify(filtered_student)


@bp.route('/api/admin/student/<int:student_id>')
def api_admin_student(student_id):
    """管理员学生相关接口。 参数: student_id"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    student = models.get_student_by_id(student_id)
    if not student:
        return jsonify({'error': '学生不存在'}), 404
    student['age'] = utils.calculate_age(student.get('birth_date'))
    return jsonify(student)


@bp.route('/api/admin/student/<int:student_id>', methods=['POST'])
def api_admin_student_update(student_id):
    """管理员学生更新相关接口。 参数: student_id"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    data = request.get_json()
    if not data:
        return jsonify({'error': '无效数据'}), 400

    required_fields = [
        ('name', '学生姓名'),
        ('id_card', '身份证号'),
        ('ethnicity', '民族'),
        ('major', '专业'),
        ('admin_class', '行政班级'),
        ('admission_type', '入学形式')
    ]

    for field_key, field_name in required_fields:
        field_value = data.get(field_key, '').strip()
        if not field_value:
            return jsonify({'error': f'请填写必填字段：{field_name}'}), 400

    id_card = data.get('id_card', '').strip()
    if id_card:
        check = utils.validate_id_card_detailed(id_card)
        if not check['valid']:
            return jsonify({'error': check['message']}), 400
        if not data.get('gender', '').strip():
            gender_from_id = check.get('gender')
            if gender_from_id:
                data['gender'] = gender_from_id

    parent_id_card = data.get('parent_id_card', '').strip()
    if parent_id_card and parent_id_card != id_card:
        pcheck = utils.validate_id_card_detailed(parent_id_card)
        if not pcheck['valid']:
            return jsonify({'error': '家长' + pcheck['message']}), 400

    parent_phone = data.get('parent_phone', '').strip()
    if parent_phone and not utils.validate_phone(parent_phone):
        return jsonify({'error': '家长手机号格式不正确'}), 400

    student_phone = data.get('student_phone', '').strip()
    if student_phone and not utils.validate_phone(student_phone):
        return jsonify({'error': '学生手机号格式不正确'}), 400

    student_data = {
        'student_no': data.get('student_no', '').strip(),
        'national_no': data.get('national_no', '').strip(),
        'id_card': id_card,
        'name': data.get('name', '').strip(),
        'gender': data.get('gender', ''),
        'ethnicity': data.get('ethnicity', '').strip(),
        'address': data.get('address', '').strip(),
        'major': data.get('major', '').strip(),
        'admin_class': data.get('admin_class', '').strip(),
        'parent_name': data.get('parent_name', '').strip(),
        'parent_id_card': data.get('parent_id_card', '').strip(),
        'parent_phone': parent_phone,
        'student_phone': student_phone,
        'photo_url': data.get('photo_url', '').strip(),
        'issue_authority': data.get('issue_authority', '').strip(),
        'valid_start': data.get('valid_start', '').strip(),
        'valid_end': data.get('valid_end', '').strip(),
        'major_direction': data.get('major_direction', '').strip(),
        'household_type': data.get('household_type', '').strip(),
        'junior_school': data.get('junior_school', '').strip(),
        'admission_type': data.get('admission_type', '').strip(),
        'junior_graduation_time': data.get('junior_graduation_time', '').strip(),
        'change_info': data.get('change_info', '').strip(),
        'remark': data.get('remark', '').strip(),
        'grade': data.get('grade', '').strip(),
        'virtual_class': data.get('virtual_class', '').strip(),
        'name_pinyin': pinyin_utils.name_to_pinyin(data.get('name', '').strip())
    }
    if 'major' in data:
        frontend_school = data.get('school')
        if frontend_school and int(frontend_school) > 0:
            student_data['school'] = int(frontend_school)
        else:
            student_data['school'] = utils._derive_school(data.get('major', ''))[0]

    if id_card:
        birth_date = utils.extract_birth_date_from_id_card(id_card)
        student_data['birth_date'] = birth_date

    try:
        models.update_student(student_id, student_data)
        models.update_sync_timestamp()
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@bp.route('/api/admin/student/add', methods=['POST'])
def api_admin_student_add():
    """管理员学生add相关接口。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    data = request.get_json()
    if not data:
        return jsonify({'error': '无效数据'}), 400

    required_fields = [
        ('name', '学生姓名'),
        ('id_card', '身份证号'),
        ('ethnicity', '民族'),
        ('major', '专业'),
        ('admin_class', '行政班级'),
        ('admission_type', '入学形式')
    ]

    for field_key, field_name in required_fields:
        field_value = data.get(field_key, '').strip()
        if not field_value:
            return jsonify({'error': f'请填写必填字段：{field_name}'}), 400

    id_card = data.get('id_card', '').strip()
    if id_card:
        check = utils.validate_id_card_detailed(id_card)
        if not check['valid']:
            return jsonify({'error': check['message']}), 400
        if not data.get('gender', '').strip():
            gender_from_id = check.get('gender')
            if gender_from_id:
                data['gender'] = gender_from_id

    parent_id_card = data.get('parent_id_card', '').strip()
    if parent_id_card and parent_id_card != id_card:
        pcheck = utils.validate_id_card_detailed(parent_id_card)
        if not pcheck['valid']:
            return jsonify({'error': '家长' + pcheck['message']}), 400

    student_data = {
        'student_no': data.get('student_no', '').strip(),
        'national_no': data.get('national_no', '').strip(),
        'id_card': id_card,
        'name': data.get('name', '').strip(),
        'gender': data.get('gender', ''),
        'ethnicity': data.get('ethnicity', '').strip(),
        'address': data.get('address', '').strip(),
        'major': data.get('major', '').strip(),
        'admin_class': data.get('admin_class', '').strip(),
        'virtual_class': data.get('virtual_class', '').strip(),
        'parent_name': data.get('parent_name', '').strip(),
        'parent_id_card': data.get('parent_id_card', '').strip(),
        'parent_phone': data.get('parent_phone', '').strip(),
        'student_phone': data.get('student_phone', '').strip(),
        'grade': data.get('grade', '').strip(),
        'school': int(data.get('school', 0)) or utils._derive_school(data.get('major', ''))[0],
        'registration_status': data.get('registration_status', '未注册'),
        'attendance_status': data.get('attendance_status', '在校'),
        'photo_url': data.get('photo_url', '').strip(),
        'issue_authority': data.get('issue_authority', '').strip(),
        'valid_start': data.get('valid_start', '').strip(),
        'valid_end': data.get('valid_end', '').strip(),
        'major_direction': data.get('major_direction', '').strip(),
        'household_type': data.get('household_type', '').strip(),
        'junior_school': data.get('junior_school', '').strip(),
        'admission_type': data.get('admission_type', '').strip(),
        'junior_graduation_time': data.get('junior_graduation_time', '').strip(),
        'change_info': data.get('change_info', '').strip(),
        'remark': data.get('remark', '').strip(),
        'name_pinyin': pinyin_utils.name_to_pinyin(data.get('name', '').strip())
    }

    year = _get_current_enrollment_year()
    if year:
        student_data['enrollment_year'] = year

    if id_card:
        birth_date = utils.extract_birth_date_from_id_card(id_card)
        student_data['birth_date'] = birth_date

    try:
        student_id = models.create_student(student_data)
        models.update_sync_timestamp()
        return jsonify({'success': True, 'student_id': student_id})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@bp.route('/api/read_id_card', methods=['POST'])
def api_read_id_card():
    """读取id身份证。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    try:
        reader = get_id_card_reader(photos_dir=config.PHOTOS_DIR)
        year = _get_current_enrollment_year()
        success, result = reader.read_id_card(year=year)

        if not success:
            return jsonify({'success': False, 'message': result}), 500

        return jsonify({
            'success': True,
            'data': {
                'name': result.get('name', ''),
                'gender': result.get('gender', ''),
                'nation': result.get('nation', ''),
                'birthday': result.get('birthday', ''),
                'address': result.get('address', ''),
                'id_card': result.get('id_card', ''),
                'issue_authority': result.get('issue_authority', ''),
                'valid_start': result.get('valid_start', ''),
                'valid_end': result.get('valid_end', ''),
                'photo_url': result.get('photo_url', ''),
                'photo_base64': result.get('photo_base64', ''),
            }
        })

    except Exception as e:
        return jsonify({'success': False, 'message': f'读卡异常: {str(e)}'}), 500


@bp.route('/api/admin/card-reader-config', methods=['GET', 'POST'])
@auth.admin_required
def api_card_reader_config():
    """读卡读卡器配置。"""
    if request.method == 'GET':
        return jsonify({'url': utils._get_card_reader_ws_url()})
    data = request.get_json()
    url = (data or {}).get('url', '').strip()
    if not url.startswith('ws://'):
        return jsonify({'success': False, 'message': '地址必须以 ws:// 开头'}), 400
    utils._set_card_reader_ws_url(url)
    return jsonify({'success': True, 'url': url, 'message': '读卡器地址已更新，刷新页面生效'})


@bp.route('/api/photos/<id_card>')
@auth.login_required
def api_photo_by_id_card(id_card):
    """照片byid身份证相关接口。 参数: id_card"""
    safe_id = ''.join(c for c in str(id_card) if c.isalnum())
    photos_dir = config.PHOTOS_DIR

    def _send(photo_path, ext):
        """发送 HTTP 响应的辅助函数。"""
        mime = 'image/jpeg' if ext in ['.jpg', '.jpeg'] else 'image/' + ext[1:]
        return send_file(photo_path, mimetype=mime)

    def _search_in_dir(search_dir):
        """在指定目录中搜索文件。"""
        for ext in ['.jpg', '.jpeg', '.png', '.bmp']:
            photo_path = os.path.join(search_dir, f'{safe_id}{ext}')
            if os.path.exists(photo_path):
                return _send(photo_path, ext)
        return None

    year = _get_current_enrollment_year()

    if year:
        result = _search_in_dir(os.path.join(photos_dir, str(year)))
        if result:
            return result
    else:
        # 全校视图：遍历所有年份子目录查找照片
        try:
            for entry in os.scandir(photos_dir):
                if entry.is_dir():
                    result = _search_in_dir(entry.path)
                    if result:
                        return result
        except OSError:
            pass

    # 兜底：检查根目录（兼容旧数据）
    result = _search_in_dir(photos_dir)
    if result:
        return result

    return jsonify({'error': '照片不存在'}), 404


@bp.route('/api/photos/save', methods=['POST'])
def api_save_photo():
    """保存照片相关接口。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    data = request.get_json()
    id_card = data.get('id_card')
    photo_base64 = data.get('photo_base64')

    if not id_card or not photo_base64:
        return jsonify({'success': False, 'message': '身份证号和照片数据不能为空'}), 400

    safe_id = ''.join(c for c in str(id_card) if c.isalnum())
    if not safe_id:
        return jsonify({'success': False, 'message': '身份证号格式不正确'}), 400

    try:
        if photo_base64.startswith('data:image/'):
            match = re.match(r'data:image/\w+;base64,(.+)', photo_base64)
            if match:
                photo_base64 = match.group(1)

        photo_data = base64.b64decode(photo_base64)
    except Exception as e:
        return jsonify({'success': False, 'message': f'照片数据解码失败: {str(e)}'}), 400

    if len(photo_data) >= 2:
        if photo_data[:2] == b'\xff\xd8':
            ext = '.jpg'
        elif photo_data[:2] == b'BM':
            ext = '.bmp'
        else:
            ext = '.bmp'
    else:
        ext = '.bmp'

    year = _get_current_enrollment_year()
    if year is None:
        # 全校视图：从数据库查找学生的入学年份
        try:
            row = models.get_db().execute(
                'SELECT enrollment_year FROM students WHERE id_card = ?', (id_card,)
            ).fetchone()
            if row and row['enrollment_year']:
                year = int(row['enrollment_year'])
            else:
                year = datetime.now().year
        except Exception:
            year = datetime.now().year
    photos_dir = os.path.join(config.PHOTOS_DIR, str(year))
    os.makedirs(photos_dir, exist_ok=True)

    filename = f'{safe_id}{ext}'
    filepath = os.path.join(photos_dir, filename)

    if os.path.exists(filepath):
        os.remove(filepath)

    with open(filepath, 'wb') as f:
        f.write(photo_data)

    photo_url = f'/api/photos/{id_card}'
    return jsonify({
        'success': True,
        'message': '照片保存成功',
        'photo_url': photo_url,
        'filename': filename
    })


@bp.route('/api/photos/count', methods=['GET'])
def api_photos_count():
    """照片数量。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    try:
        photos_dir = config.PHOTOS_DIR
        if not os.path.exists(photos_dir):
            return jsonify({'success': True, 'count': 0})

        year = _get_current_enrollment_year()

        count = 0
        if year:
            year_dir = os.path.join(photos_dir, str(year))
            if os.path.exists(year_dir):
                for f in os.listdir(year_dir):
                    if f.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp')):
                        count += 1

        if count == 0:
            for f in os.listdir(photos_dir):
                fp = os.path.join(photos_dir, f)
                if os.path.isfile(fp) and f.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp')):
                    count += 1

        return jsonify({'success': True, 'count': count})
    except Exception as e:
        return jsonify({'success': False, 'message': f'统计照片失败: {str(e)}'}), 500


@bp.route('/api/photos/delete/<id_card>', methods=['DELETE'])
def api_photos_delete(id_card):
    """照片删除。 参数: id_card"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    safe_id = ''.join(c for c in str(id_card) if c.isalnum())
    photos_dir = config.PHOTOS_DIR

    deleted = False

    year = _get_current_enrollment_year()
    if year:
        year_dir = os.path.join(photos_dir, str(year))
        for ext in ['.jpg', '.jpeg', '.png', '.bmp']:
            photo_path = os.path.join(year_dir, f'{safe_id}{ext}')
            if os.path.exists(photo_path):
                os.remove(photo_path)
                deleted = True

    for ext in ['.jpg', '.jpeg', '.png', '.bmp']:
        photo_path = os.path.join(photos_dir, f'{safe_id}{ext}')
        if os.path.exists(photo_path):
            os.remove(photo_path)
            deleted = True

    if deleted:
        return jsonify({'success': True, 'message': '照片已删除'})
    return jsonify({'success': False, 'message': '照片不存在'}), 404


@bp.route('/api/photos/export/zip', methods=['GET'])
@auth.admin_required
def api_export_photos_zip():
    """导出照片zip。"""
    photos_dir = config.PHOTOS_DIR
    if not os.path.exists(photos_dir):
        return jsonify({'success': False, 'message': '照片目录不存在'}), 404

    year = _get_current_enrollment_year()
    search_dirs = []

    if year:
        year_dir = os.path.join(photos_dir, str(year))
        if os.path.exists(year_dir):
            search_dirs.append((year_dir, f'{year}/'))

    search_dirs.append((photos_dir, ''))

    output = BytesIO()
    count = 0
    with zipfile.ZipFile(output, 'w', zipfile.ZIP_DEFLATED) as zf:
        for search_dir, arc_prefix in search_dirs:
            for f in os.listdir(search_dir):
                fp = os.path.join(search_dir, f)
                if os.path.isfile(fp) and f.lower().endswith(('.jpg', '.jpeg', '.png', '.bmp')):
                    zf.write(fp, arc_prefix + f)
                    count += 1

    if count == 0:
        return jsonify({'success': False, 'message': '没有可导出的照片'}), 404

    output.seek(0)
    return send_file(output, mimetype='application/zip',
                     download_name=f'学生照片_{datetime.now().strftime("%Y%m%d_%H%M%S")}.zip',
                     as_attachment=True)


@bp.route('/api/students/check_id_card')
def api_check_id_card():
    """检查id身份证。"""
    auth_error = _require_login_json()
    if auth_error:
        return auth_error
    id_card = (request.args.get('id_card') or '').strip()
    exclude_id = request.args.get('exclude_id', type=int)

    if not id_card:
        return jsonify({'success': False, 'message': '身份证号不能为空'}), 400

    check = utils.validate_id_card_detailed(id_card)
    if not check['valid']:
        return jsonify({'success': True, 'valid': False, 'message': check['message']})

    existing = models.get_student_by_id_card(id_card)
    exists = existing is not None and (not exclude_id or existing.get('id') != exclude_id)

    return jsonify({
        'success': True,
        'valid': True,
        'exists': exists,
        'birth_date': check.get('birth_date'),
        'gender': check.get('gender'),
        'region': check.get('region'),
        'message': '该身份证号已存在' if exists else '身份证号可用'
    })


@bp.route('/api/dicts', methods=['GET'])
def api_get_dicts():
    """获取字典。"""
    auth_error = _require_login_json()
    if auth_error:
        return auth_error
    try:
        dict_type = request.args.get('type', '')
        search = request.args.get('search', '').strip()
        rows = models.get_dicts(
            dict_type if dict_type else None,
            search if search else None
        )
        grouped = models.get_grouped_dicts()
        return jsonify({'success': True, 'dicts': rows, 'grouped_dicts': grouped})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@bp.route('/api/dicts/types', methods=['GET'])
def api_get_dict_types():
    """获取字典types。"""
    auth_error = _require_login_json()
    if auth_error:
        return auth_error
    try:
        rows = models.get_dicts()
        types = sorted(set(r['dict_type'] for r in rows))
        return jsonify({'success': True, 'types': types})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@bp.route('/api/dicts', methods=['POST'])
def api_create_dict():
    """创建字典。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    try:
        data = request.get_json()
        dict_type = (data.get('dict_type') or '').strip()
        dict_value = (data.get('dict_value') or '').strip()
        if not dict_type or not dict_value:
            return jsonify({'success': False, 'message': '字典类型和值不能为空'}), 400
        ok = models.add_dict(dict_type, dict_value)
        if not ok:
            return jsonify({'success': False, 'message': '该字典项已存在'}), 400
        return jsonify({'success': True, 'message': '字典项创建成功'})
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@bp.route('/api/dicts/<int:dict_id>', methods=['DELETE'])
def api_delete_dict(dict_id):
    """删除字典。 参数: dict_id"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    try:
        ok = models.delete_dict(dict_id)
        if ok:
            return jsonify({'success': True, 'message': '字典项删除成功'})
        return jsonify({'success': False, 'message': '字典项不存在'}), 404
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@bp.route('/api/dicts/batch', methods=['POST'])
def api_batch_import_dicts():
    """批量导入字典。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    try:
        file = request.files.get('file')
        if not file:
            return jsonify({'success': False, 'message': '请选择文件'}), 400

        wb = load_workbook(file)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))

        imported = 0
        skipped = 0
        for row in rows:
            if not row or not row[0] or not row[1]:
                continue
            dict_type = str(row[0]).strip()
            dict_value = str(row[1]).strip()
            if not dict_type or not dict_value:
                continue
            if models.add_dict(dict_type, dict_value):
                imported += 1
            else:
                skipped += 1

        return jsonify({'success': True, 'message': f'导入完成，成功{imported}条，跳过{skipped}条'})
    except Exception as e:
        return jsonify({'success': False, 'message': f'导入失败: {str(e)}'}), 500


@bp.route('/api/dicts/export', methods=['GET'])
@auth.admin_required
def api_export_dicts():
    """导出字典。"""
    try:
        rows = models.get_dicts()
        if not rows:
            return jsonify({'success': False, 'message': '没有可导出的数据'}), 404

        wb = Workbook()
        ws = wb.active
        ws.title = '字典数据'
        ws['A1'] = 'dict_type'
        ws['B1'] = 'dict_value'
        ws['A1'].font = Font(bold=True, color='FFFFFF')
        ws['A1'].fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
        ws['B1'].font = Font(bold=True, color='FFFFFF')
        ws['B1'].fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 30

        for i, r in enumerate(rows, 2):
            ws.cell(row=i, column=1, value=r['dict_type'])
            ws.cell(row=i, column=2, value=r['dict_value'])

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         download_name=f'字典数据_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx',
                         as_attachment=True)
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@bp.route('/api/dicts/template', methods=['GET'])
@auth.admin_required
def api_dict_template():
    """字典模板相关接口。"""
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = '字典数据'

        sample_data = [
            ('nation', '汉族'),
            ('major', '护理'),
            ('major_direction', '升学'),
            ('household_type', '城市'),
            ('household_type', '县镇'),
            ('household_type', '农村'),
            ('junior_school', '测试中学'),
            ('admission_type', '普通入学'),
            ('admission_type', '转学生'),
        ]

        ws['A1'] = 'dict_type'
        ws['B1'] = 'dict_value'
        for cell in ['A1', 'B1']:
            ws[cell].font = Font(bold=True, color='FFFFFF')
            ws[cell].fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
            ws[cell].alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 30

        for i, (t, v) in enumerate(sample_data, 2):
            ws.cell(row=i, column=1, value=t)
            ws.cell(row=i, column=2, value=v)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         download_name='字典导入模板.xlsx', as_attachment=True)
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@bp.route('/api/school-types', methods=['GET'])
def api_get_school_types():
    """获取学校类别列表（含停用项）。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    try:
        rows = models.get_school_types()
        return jsonify({
            'success': True,
            'school_types': rows,
            'suggested_code': models.get_next_school_type_code()
        })
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@bp.route('/api/school-types', methods=['POST'])
def api_save_school_type():
    """新增/改名/停用启用学校类别（action: create | update）。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    try:
        data = request.get_json(silent=True) or {}
        action = (data.get('action') or 'create').strip()
        code = data.get('code')

        if action == 'create':
            created = models.create_school_type(code, data.get('name'))
            # 新类别自动追加到 admin 可见范围，避免管理员看不到新类别数据
            models.append_code_to_admin_visibility(created['code'])
            return jsonify({'success': True, 'message': '学校类别创建成功'})

        if action == 'update':
            is_active = data.get('is_active')
            if is_active is not None:
                if is_active in (True, 1, '1'):
                    is_active = 1
                elif is_active in (False, 0, '0'):
                    is_active = 0
                else:
                    return jsonify({'success': False, 'message': '启用状态参数无效'}), 400
            ok = models.update_school_type(
                code,
                name=data.get('name'),
                is_active=is_active
            )
            if not ok:
                return jsonify({'success': False, 'message': '学校类别不存在'}), 404
            return jsonify({'success': True, 'message': '学校类别更新成功'})

        return jsonify({'success': False, 'message': '不支持的操作'}), 400
    except ValueError as e:
        return jsonify({'success': False, 'message': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@bp.route('/api/school-types/<int:code>', methods=['DELETE'])
def api_delete_school_type(code):
    """删除学校类别（仅未被学生或用户权限引用时允许）。"""
    auth_error = _require_admin_json()
    if auth_error:
        return auth_error
    try:
        ok = models.delete_school_type(code)
        if not ok:
            return jsonify({'success': False, 'message': '学校类别不存在'}), 404
        return jsonify({'success': True, 'message': '学校类别删除成功'})
    except ValueError as e:
        return jsonify({'success': False, 'message': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'message': str(e)}), 500


@bp.route('/api/stats')
def api_stats():
    """统计。"""
    auth_error = _require_login_json()
    if auth_error:
        return auth_error
    filters = auth.apply_permission_filters({})
    stats = models.get_stats(filters)
    return jsonify(stats)

@bp.route('/api/export')
@auth.login_required
def api_export():
    """导出。"""
    filters = utils.parse_filters_from_request(request)
    filters = auth.apply_permission_filters(filters)
    
    result = models.get_students(filters, per_page=10000)
    students = result['students']
    
    visible_fields = auth.get_visible_fields()
    export_fields = visible_fields if visible_fields != ['*'] else None
    
    for student in students:
        student['age'] = utils.calculate_age(student.get('birth_date'))
    
    wb = utils.export_students_to_excel(students, export_fields)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    timestamp = datetime.now().strftime('%Y%m%d_%H%M')
    filename = f'学生信息导出_{timestamp}.xlsx'
    
    return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=filename)

@bp.route('/api/filters/options')
@auth.login_required
def api_filters_options():
    """筛选options。"""
    school = request.args.get('school', '')
    grade = request.args.get('grade', '')

    base_school_visible = auth.effective_school_visible(school)

    filter_values = models.get_filter_values(school_visible=base_school_visible, grade_visible='*')

    result = {
        'grades': filter_values['grades'],
        'majors': filter_values['majors'],
        'classes': filter_values['classes'],
        'virtual_classes': filter_values['virtual_classes'],
        'ethnicities': filter_values['ethnicities']
    }

    if grade:
        result['classes'] = models.get_classes_by_grade(grade, base_school_visible)

    return jsonify(result)

@bp.route('/api/filters/classes')
@auth.login_required
def api_filters_classes():
    """筛选班级。"""
    grade = request.args.get('grade', '')
    school = request.args.get('school', '')

    school_visible = auth.effective_school_visible(school)

    if grade:
        classes = models.get_classes_by_grade(grade, school_visible)
    else:
        filter_values = models.get_filter_values(school_visible)
        classes = filter_values['classes']
    
    return jsonify({'classes': classes})

@bp.route('/api/filters/majors')
@auth.login_required
def api_filters_majors():
    """筛选专业。"""
    school = request.args.get('school', '')
    school_visible = auth.effective_school_visible(school)
    grade_visible = session.get('grade_visible') or '*'
    filter_values = models.get_filter_values(school_visible, grade_visible)
    return jsonify({'majors': filter_values['majors']})

@bp.route('/api/filters/virtual-classes')
@auth.login_required
def api_filters_virtual_classes():
    """筛选虚拟班级。"""
    school = request.args.get('school', '')
    school_visible = auth.effective_school_visible(school)
    grade_visible = session.get('grade_visible') or '*'
    filter_values = models.get_filter_values(school_visible, grade_visible)
    return jsonify({'virtual_classes': filter_values['virtual_classes']})

@bp.route('/api/grades')
@auth.login_required
def api_grades():
    """年级。"""
    school = request.args.get('school', '')
    school_visible = auth.effective_school_visible(school)
    filter_values = models.get_filter_values(school_visible, grade_visible='*')
    return jsonify({'grades': filter_values['grades']})

@bp.route('/api/attendance-dates')
@auth.login_required
def api_attendance_dates():
    """出勤日期。"""
    dates = models.get_available_attendance_dates()
    return jsonify({'dates': dates})

@bp.route('/api/today-attendance')
@auth.login_required
def api_today_attendance():
    """今日出勤。"""
    school = request.args.get('school', '')
    registration_status = request.args.get('registration_status', '')
    
    filters = {}
    if school:
        filters['school'] = int(school)
    if registration_status in ('正式学籍', '未注册'):
        filters['registration_status'] = registration_status
    
    year = _get_current_enrollment_year()
    if year is not None:
        filters['enrollment_year'] = year
    
    extra_clauses = []
    extra_params = []
    
    if filters.get('school'):
        extra_clauses.append('s.school = ?')
        extra_params.append(filters['school'])
    if filters.get('registration_status'):
        extra_clauses.append('s.registration_status = ?')
        extra_params.append(filters['registration_status'])
    if filters.get('enrollment_year'):
        extra_clauses.append('s.enrollment_year = ?')
        extra_params.append(filters['enrollment_year'])
    
    extra_where = (' AND ' + ' AND '.join(extra_clauses)) if extra_clauses else ''
    
    today = datetime.now().strftime('%Y-%m-%d')
    
    conn = models.get_db()
    cursor = conn.cursor()
    
    cursor.execute(f'''
        SELECT ar.status, COUNT(DISTINCT ar.id_card) as count
        FROM attendance_records ar
        INNER JOIN students s ON ar.id_card = s.id_card AND s.is_deleted = 0 AND s.attendance_status NOT IN ('非在校', '退学办结') AND (s.registration_status IS NULL OR s.registration_status != '毕业')
        WHERE ar.attendance_date = ?{extra_where}
        GROUP BY ar.status
    ''', [today] + extra_params)
    rows = {row['status']: row['count'] for row in [dict(r) for r in cursor.fetchall()]}

    present = rows.get('出勤', 0)
    absent = rows.get('缺勤', 0)
    total = present + absent

    cursor.execute(f'''
        SELECT s.admin_class,
               COUNT(DISTINCT s.id) as total_students,
               SUM(CASE WHEN ar.status = '出勤' THEN 1 ELSE 0 END) as present,
               SUM(CASE WHEN ar.status = '缺勤' THEN 1 ELSE 0 END) as absent
        FROM students s
        LEFT JOIN attendance_records ar ON s.id_card = ar.id_card AND ar.attendance_date = ?
        WHERE s.is_deleted = 0 AND s.attendance_status NOT IN ('非在校', '退学办结') AND (s.registration_status IS NULL OR s.registration_status != '毕业'){extra_where} AND s.admin_class IS NOT NULL AND s.admin_class != ''
        GROUP BY s.admin_class
    ''', [today] + extra_params)
    by_class = [dict(row) for row in cursor.fetchall()]
    
    import re
    from utils import natural_key as _natural_key  # P2-2: 统一自然排序键

    def _rate_sort_key(row):
        name = row.get('admin_class') or ''
        m = re.match(r'^(\d{2})', name)
        year = int(m.group(1)) if m else 0
        rest = name[2:] if m else name
        return (year, _natural_key(rest))
    
    by_class.sort(key=_rate_sort_key)
    
    conn.close()
    
    return jsonify({
        'present': present,
        'absent': absent,
        'total': total,
        'by_class': [{
            'admin_class': r['admin_class'],
            'total_students': r['total_students'],
            'present': r['present'] or 0,
            'absent': r['absent'] or 0
        } for r in by_class],
        'updated_at': datetime.now().strftime('%H:%M:%S')
    })

@bp.route('/api/class-comparison')
@auth.login_required
def api_class_comparison():
    """班级对比相关接口。"""
    school = request.args.get('school', '')
    registration_status = request.args.get('registration_status', '')

    filters = {}
    if school:
        filters['school'] = int(school)
    if registration_status:
        filters['registration_status'] = registration_status

    filters = auth.apply_permission_filters(filters)

    data = models.get_class_comparison(
        school=filters.get('school'),
        registration_status=filters.get('registration_status'),
        grade=filters.get('grade')
    )
    return jsonify(data)

@bp.route('/api/class-comparison-detail')
@auth.login_required
def api_class_comparison_detail():
    """班级对比详情相关接口。"""
    school = request.args.get('school', '')
    registration_status = request.args.get('registration_status', '')

    filters = {}
    if school:
        filters['school'] = int(school)
    if registration_status:
        filters['registration_status'] = registration_status

    filters = auth.apply_permission_filters(filters)

    data = models.get_class_comparison_detail(
        school=filters.get('school'),
        registration_status=filters.get('registration_status'),
        grade=filters.get('grade')
    )
    return jsonify(data)
