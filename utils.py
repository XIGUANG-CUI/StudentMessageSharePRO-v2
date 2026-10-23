import re
import os
import random
import string
import socket
import ipaddress
import platform
import logging
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
import config
import models

logger = logging.getLogger(__name__)


def natural_key(name):
    """自然排序键：将字符串中的数字部分转为整数，实现 '班级2' < '班级10' 的自然顺序。
    用于班级/专业等含数字名称的排序。None/空字符串返回空列表（健壮处理）。
    P2-2: 统一替换 models.py / blueprints/api.py 中 4 处重复定义。"""
    name = name or ''
    parts = re.split(r'(\d+)', name)
    result = []
    for i, part in enumerate(parts):
        if i % 2 == 1:
            result.append(int(part))
        else:
            result.append(part)
    return result


def extract_birth_date_from_id_card(id_card):
    """从 18 位身份证号提取出生日期（YYYY-MM-DD）。 参数: id_card"""
    if not id_card or len(id_card) != 18:
        return None
    
    try:
        birth_date_str = id_card[6:14]
        birth_date = datetime.strptime(birth_date_str, '%Y%m%d').date()
        return birth_date.strftime('%Y-%m-%d')
    except (ValueError, TypeError):
        return None

def calculate_age(birth_date):
    """根据出生日期计算年龄。 参数: birth_date"""
    if not birth_date:
        return None
    
    try:
        if isinstance(birth_date, str):
            birth_date = datetime.strptime(birth_date, '%Y-%m-%d').date()
        
        today = datetime.now().date()
        years = today.year - birth_date.year
        months = today.month - birth_date.month
        
        if months < 0 or (months == 0 and today.day < birth_date.day):
            years -= 1
            months += 12
        
        if today.day < birth_date.day:
            months -= 1
        
        return f'{years}岁{months}个月'
    except (ValueError, TypeError):
        return None

def validate_id_card(id_card):
    """校验id身份证数据。 参数: id_card"""
    if not id_card:
        return False

    id_card = str(id_card).strip().upper()

    if len(id_card) == 15:
        return True

    if len(id_card) != 18:
        return False

    if not id_card[:-1].isdigit():
        return False

    last_char = id_card[-1]
    if not (last_char.isdigit() or last_char == 'X'):
        return False

    weights = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    check_codes = ['1', '0', 'X', '9', '8', '7', '6', '5', '4', '3', '2']

    total = 0
    for i in range(17):
        total += int(id_card[i]) * weights[i]

    expected = check_codes[total % 11]
    return last_char == expected


_ID_CARD_REGION = {
    '11': '北京', '12': '天津', '13': '河北', '14': '山西', '15': '内蒙古',
    '21': '辽宁', '22': '吉林', '23': '黑龙江',
    '31': '上海', '32': '江苏', '33': '浙江', '34': '安徽', '35': '福建',
    '36': '江西', '37': '山东',
    '41': '河南', '42': '湖北', '43': '湖南', '44': '广东', '45': '广西',
    '46': '海南',
    '50': '重庆', '51': '四川', '52': '贵州', '53': '云南', '54': '西藏',
    '61': '陕西', '62': '甘肃', '63': '青海', '64': '宁夏', '65': '新疆',
    '71': '台湾', '81': '香港', '82': '澳门',
    '91': '国外',
}


def validate_id_card_detailed(id_card):
    """校验id身份证detailed数据。 参数: id_card"""
    result = {'valid': True, 'error_code': None, 'message': '', 'birth_date': None, 'gender': None, 'region': None}

    if not id_card:
        return _make_error(result, 'EMPTY', '身份证号码不能为空')

    id_card = str(id_card).strip().upper()

    if len(id_card) == 0:
        return _make_error(result, 'EMPTY', '身份证号码不能为空')

    if len(id_card) != 18:
        return _make_error(result, 'INVALID_LENGTH', '身份证号码长度不正确，必须为18位')

    if not id_card[:-1].isdigit():
        return _make_error(result, 'INVALID_FORMAT', '身份证号码前17位必须为数字')

    last_char = id_card[-1]
    if not (last_char.isdigit() or last_char == 'X'):
        return _make_error(result, 'INVALID_FORMAT', '身份证号码第18位必须为数字或大写字母X')

    region_code = id_card[:2]
    if region_code not in _ID_CARD_REGION:
        return _make_error(result, 'INVALID_REGION', '身份证号码地区码无效')

    birth_str = id_card[6:14]
    try:
        birth_date = datetime.strptime(birth_str, '%Y%m%d').date()
        if birth_date > datetime.now().date():
            return _make_error(result, 'INVALID_BIRTH', '出生日期不能是未来日期')
        if birth_date.year < 1900:
            return _make_error(result, 'INVALID_BIRTH', '出生日期无效')
    except ValueError:
        return _make_error(result, 'INVALID_BIRTH', '身份证号码出生日期无效')

    gender_digit = int(id_card[16])
    result['gender'] = '男' if gender_digit % 2 == 1 else '女'

    weights = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
    check_codes = ['1', '0', 'X', '9', '8', '7', '6', '5', '4', '3', '2']

    total = 0
    for i in range(17):
        total += int(id_card[i]) * weights[i]

    expected = check_codes[total % 11]
    if last_char != expected:
        return _make_error(result, 'CHECKSUM_FAIL', '身份证号码校验位验证失败，请检查输入是否正确')

    result['region'] = _ID_CARD_REGION[region_code]
    result['birth_date'] = birth_str

    return result


def _make_error(result, code, msg):
    """构造校验错误响应对象。 参数: result, code, msg"""
    result['valid'] = False
    result['error_code'] = code
    result['message'] = msg
    return result


_VALID_PHONE_PREFIXES = {
    '134', '135', '136', '137', '138', '139',
    '147', '148',
    '150', '151', '152', '157', '158', '159',
    '165',
    '172', '178',
    '182', '183', '184', '187', '188',
    '195', '198',
    '130', '131', '132',
    '145', '146',
    '155', '156',
    '166',
    '171', '175', '176',
    '185', '186',
    '196',
    '133',
    '149',
    '153',
    '162',
    '173', '174', '177',
    '180', '181', '189',
    '190', '191', '193', '199',
    '192',
    '170',
}


def validate_phone(phone):
    """校验电话数据。 参数: phone"""
    if not phone:
        return True

    phone = str(phone).strip()
    phone = phone.replace(' ', '').replace('-', '')

    if not re.match(r'^1\d{10}$', phone):
        return False

    prefix = phone[:3]

    if prefix in _VALID_PHONE_PREFIXES:
        return True

    second_digit = phone[1]
    if '3' <= second_digit <= '9':
        return True

    return False


def generate_random_password(length=12):
    """生成随机密码。 参数: length"""
    chars = string.ascii_letters + string.digits + '!@#$%^&*'
    password = ''.join(random.choice(chars) for _ in range(length))
    return password

def allowed_file(filename):
    """检查文件扩展名是否在允许列表内。 参数: filename"""
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in config.ALLOWED_EXTENSIONS

def generate_unique_filename(original_filename):
    """生成唯一文件名。 参数: original_filename"""
    ext = original_filename.rsplit('.', 1)[1].lower()
    random_str = ''.join(random.choices(string.ascii_letters + string.digits, k=8))
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    return f'{timestamp}_{random_str}.{ext}'

def create_import_template():
    """创建导入模板。"""
    wb = Workbook()
    ws = wb.active
    ws.title = '导入模板'
    
    headers = [
        '学号', '国网学籍号', '学生身份证号', '学生姓名', '性别', 
        '民族', '住址', '专业', '行政班级', '虚拟班级', 
        '家长姓名', '家长身份证号', '家长电话', '学生电话', 
        '年级', '学校类型', '学籍状态', '在校状态',
        '签发机关', '有效期开始', '有效期结束', '专业方向',
        '户口性质', '初中毕业学校', '入学形式', '初中毕业时间',
        '异动信息', '备注', '入学年份'
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
    
    sample_data = [
        ['2023001', 'G202300001', '110101200801151234', '张小明', '男', 
         '汉族', '北京市朝阳区XX街道XX号', '计算机应用', '计应2301班', '计应2301-1班', 
         '张建国', '110101197805201234', '13800138001', '13900139001', 
         '2023级', '1', '正式学籍', '在校',
         '北京市公安局朝阳分局', '20210601', '20210601', '就业',
         '城市', '测试第一初中', '普通入学', '202106', '', '', '2023'],
        ['2023002', 'G202300002', '310101200805205678', '李小红', '女', 
         '汉族', '上海市浦东新区XX路XX号', '数控技术', '数控2301班', '数控2301-2班', 
         '李明华', '310101197508155678', '13800138002', '13900139002', 
         '2023级', '2', '正式学籍', '在校',
         '上海市公安局浦东分局', '20220501', '20220501', '升学',
         '县镇', '测试第二初中', '转学生', '202205', '', '', '2023']
    ]
    
    for row_idx, row_data in enumerate(sample_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if row_idx == 2:
                cell.fill = PatternFill(start_color='E8F4F8', end_color='E8F4F8', fill_type='solid')
            else:
                cell.fill = PatternFill(start_color='FFF3E0', end_color='FFF3E0', fill_type='solid')
    
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 8
    ws.column_dimensions['F'].width = 10
    ws.column_dimensions['G'].width = 30
    ws.column_dimensions['H'].width = 20
    ws.column_dimensions['I'].width = 15
    ws.column_dimensions['J'].width = 15
    ws.column_dimensions['K'].width = 12
    ws.column_dimensions['L'].width = 20
    ws.column_dimensions['M'].width = 15
    ws.column_dimensions['N'].width = 15
    ws.column_dimensions['O'].width = 10
    ws.column_dimensions['P'].width = 12
    ws.column_dimensions['Q'].width = 12
    ws.column_dimensions['R'].width = 12
    ws.column_dimensions['S'].width = 20
    ws.column_dimensions['T'].width = 12
    ws.column_dimensions['U'].width = 12
    ws.column_dimensions['V'].width = 12
    ws.column_dimensions['W'].width = 10
    ws.column_dimensions['X'].width = 20
    ws.column_dimensions['Y'].width = 12
    ws.column_dimensions['Z'].width = 12
    ws.column_dimensions['AA'].width = 12
    ws.column_dimensions['AB'].width = 20
    ws.column_dimensions['AC'].width = 12
    
    ws.cell(row=1, column=30, value='填写说明：')
    ws.cell(row=2, column=30, value='学校类型：1=中职，2=技工')
    ws.cell(row=3, column=30, value='学籍状态：正式学籍/未注册')
    ws.cell(row=4, column=30, value='在校状态：在校/非在校/托管')
    ws.cell(row=5, column=30, value='入学年份：填写如2023，不填则默认0')
    ws.cell(row=6, column=30, value='新字段为选填，不填留空即可')
    
    ws.column_dimensions['AD'].width = 30
    
    return wb

def create_batch_update_template(students, attendance_date=None):
    """创建批量更新模板。 参数: students, attendance_date"""
    wb = Workbook()
    ws = wb.active
    ws.title = '批量更新模板'
    
    if attendance_date is None:
        attendance_date = datetime.now().strftime('%Y-%m-%d')

    attendance_map = {r['id_card']: r['status'] for r in models.get_attendance_records_by_date(attendance_date)}
    
    headers = ['身份证号', '姓名', '学号', '国网学籍号'] + [
        '学籍状态', '在校状态', '行政班级', '虚拟班级',
        '专业', '年级', '学生电话', '家长电话', '住址',
        '性别', '民族', '家长姓名', '家长身份证号',
        '学校类型', '专业方向', '户口性质', '初中毕业学校',
        '入学形式', '初中毕业时间', '签发机关', '有效期开始',
        '有效期结束', '异动信息', '备注', '入学年份', '今日出勤'
    ]
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
    
    for row, student in enumerate(students, 2):
        ws.cell(row=row, column=1, value=student.get('id_card', ''))
        ws.cell(row=row, column=2, value=student.get('name', ''))
        ws.cell(row=row, column=3, value=student.get('student_no', ''))
        ws.cell(row=row, column=4, value=student.get('national_no', ''))
        ws.cell(row=row, column=5, value=student.get('registration_status', ''))
        ws.cell(row=row, column=6, value=student.get('attendance_status', ''))
        ws.cell(row=row, column=7, value=student.get('admin_class', ''))
        ws.cell(row=row, column=8, value=student.get('virtual_class', ''))
        ws.cell(row=row, column=9, value=student.get('major', ''))
        ws.cell(row=row, column=10, value=student.get('grade', ''))
        ws.cell(row=row, column=11, value=student.get('student_phone', ''))
        ws.cell(row=row, column=12, value=student.get('parent_phone', ''))
        ws.cell(row=row, column=13, value=student.get('address', ''))
        ws.cell(row=row, column=14, value=student.get('gender', ''))
        ws.cell(row=row, column=15, value=student.get('ethnicity', ''))
        ws.cell(row=row, column=16, value=student.get('parent_name', ''))
        ws.cell(row=row, column=17, value=student.get('parent_id_card', ''))
        ws.cell(row=row, column=18, value=student.get('school'))
        ws.cell(row=row, column=19, value=student.get('major_direction', ''))
        ws.cell(row=row, column=20, value=student.get('household_type', ''))
        ws.cell(row=row, column=21, value=student.get('junior_school', ''))
        ws.cell(row=row, column=22, value=student.get('admission_type', ''))
        ws.cell(row=row, column=23, value=student.get('junior_graduation_time', ''))
        ws.cell(row=row, column=24, value=student.get('issue_authority', ''))
        ws.cell(row=row, column=25, value=student.get('valid_start', ''))
        ws.cell(row=row, column=26, value=student.get('valid_end', ''))
        ws.cell(row=row, column=27, value=student.get('change_info', ''))
        ws.cell(row=row, column=28, value=student.get('remark', ''))
        ey_val = student.get('enrollment_year')
        ws.cell(row=row, column=29, value=ey_val if ey_val and ey_val != 0 else None)
        
        id_card = student.get('id_card', '')
        att_status = attendance_map.get(id_card)
        if att_status == '出勤':
            ws.cell(row=row, column=30, value='出勤')
        elif att_status == '缺勤':
            ws.cell(row=row, column=30, value='缺勤')
        else:
            ws.cell(row=row, column=30, value='')
    
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 12
    for col_letter in ['C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L',
                       'M', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U', 'V',
                       'W', 'X', 'Y', 'Z', 'AA', 'AB', 'AC', 'AD']:
        ws.column_dimensions[col_letter].width = 14

    return wb

def parse_import_file(file_path):
    """解析导入文件。 参数: file_path"""
    try:
        wb = load_workbook(file_path)
        ws = wb.active
        
        headers = [cell.value for cell in ws[1]]
        
        data = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            
            row_data = dict(zip(headers, row))
            data.append(row_data)
        
        return data, None
    except Exception as e:
        return None, str(e)

def parse_batch_update_file(file_path):
    """解析批量更新文件。 参数: file_path"""
    try:
        wb = load_workbook(file_path)
        ws = wb.active
        
        headers = [cell.value for cell in ws[1]]
        
        data = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row[:2]):
                continue
            
            row_data = dict(zip(headers, row))
            data.append(row_data)
        
        return data, None
    except Exception as e:
        return None, str(e)

def _safe_str(value):
    """安全地将值转换为字符串，处理None和非字符串类型"""
    if value is None:
        return ''
    return str(value).strip() if isinstance(value, str) else str(value)

def validate_import_data(data):
    """校验导入数据数据。 参数: data"""
    results = []
    errors = []
    
    seen_student_nos = {}
    seen_national_nos = {}
    seen_id_cards = {}
    
    for idx, row in enumerate(data, 1):
        result = {
            'row': idx,
            'data': row,
            'valid': True,
            'errors': []
        }
        
        student_no = _safe_str(row.get('学号', ''))
        national_no = _safe_str(row.get('国网学籍号', ''))
        id_card = _safe_str(row.get('学生身份证号', ''))
        name = _safe_str(row.get('学生姓名', ''))
        school_type = _safe_str(row.get('学校类型', ''))
        registration_status = _safe_str(row.get('学籍状态', '正式学籍'))
        
        if not name:
            result['valid'] = False
            result['errors'].append('姓名不能为空')
        
        if id_card:
            if not validate_id_card(id_card):
                result['valid'] = False
                result['errors'].append('身份证号格式不正确')
            
            if id_card in seen_id_cards:
                result['valid'] = False
                result['errors'].append(f'身份证号与第{seen_id_cards[id_card]}行重复')
            else:
                seen_id_cards[id_card] = idx
            
            existing = models.get_student_by_id_card(id_card)
            if existing:
                result['valid'] = False
                result['errors'].append('身份证号已存在')
        else:
            result['errors'].append('身份证号不能为空')
            result['valid'] = False
        
        if student_no:
            if student_no in seen_student_nos:
                result['valid'] = False
                result['errors'].append(f'学号与第{seen_student_nos[student_no]}行重复')
            else:
                seen_student_nos[student_no] = idx
            
            existing = models.get_student_by_student_no(student_no)
            if existing:
                result['valid'] = False
                result['errors'].append('学号已存在')
        
        if national_no:
            if national_no in seen_national_nos:
                result['valid'] = False
                result['errors'].append(f'国网学籍号与第{seen_national_nos[national_no]}行重复')
            else:
                seen_national_nos[national_no] = idx
            
            existing = models.get_student_by_national_no(national_no)
            if existing:
                result['valid'] = False
                result['errors'].append('国网学籍号已存在')
        
        if school_type:
            if str(school_type) not in ['1', '2']:
                result['valid'] = False
                result['errors'].append('学校类型必须是1或2')
        else:
            result['errors'].append('学校类型不能为空')
            result['valid'] = False
        
        if registration_status not in config.REGISTRATION_STATUS:
            result['valid'] = False
            result['errors'].append(f'学籍状态必须是: {", ".join(config.REGISTRATION_STATUS)}')
        
        attendance_status = _safe_str(row.get('在校状态', '在校'))
        if attendance_status not in config.ATTENDANCE_STATUS:
            result['valid'] = False
            result['errors'].append(f'在校状态必须是: {", ".join(config.ATTENDANCE_STATUS)}')
        
        results.append(result)
    
    return results

_BATCH_UPDATE_FIELD_MAPPING = {
    '学号': 'student_no',
    '国网学籍号': 'national_no',
    '学籍状态': 'registration_status',
    '在校状态': 'attendance_status',
    '行政班级': 'admin_class',
    '虚拟班级': 'virtual_class',
    '专业': 'major',
    '年级': 'grade',
    '学生电话': 'student_phone',
    '家长电话': 'parent_phone',
    '住址': 'address',
    '性别': 'gender',
    '民族': 'ethnicity',
    '家长姓名': 'parent_name',
    '家长身份证号': 'parent_id_card',
    '学校类型': 'school',
    '专业方向': 'major_direction',
    '户口性质': 'household_type',
    '初中毕业学校': 'junior_school',
    '入学形式': 'admission_type',
    '初中毕业时间': 'junior_graduation_time',
    '签发机关': 'issue_authority',
    '有效期开始': 'valid_start',
    '有效期结束': 'valid_end',
    '异动信息': 'change_info',
    '备注': 'remark',
    '入学年份': 'enrollment_year',
    '今日出勤': '_attendance'
}


def validate_batch_update_data(data, attendance_date=None):
    """校验批量更新数据数据。 参数: data, attendance_date"""
    results = []
    
    if attendance_date is None:
        attendance_date = datetime.now().strftime('%Y-%m-%d')
    
    today_attendance = {}
    if attendance_date:
        for r in models.get_attendance_records_by_date(attendance_date):
            today_attendance[r['id_card']] = r['status']
    
    id_card_rows = {}
    
    for idx, row in enumerate(data, 1):
        result = {
            'row': idx,
            'data': row,
            'valid': True,
            'changes': [],
            'status': 'no_change'
        }
        
        id_card = _safe_str(row.get('身份证号', ''))
        name = _safe_str(row.get('姓名', ''))
        
        if not id_card:
            result['valid'] = False
            result['status'] = 'not_found'
            results.append(result)
            continue
        
        if id_card in id_card_rows:
            id_card_rows[id_card] = idx
        
        field_mapping = _BATCH_UPDATE_FIELD_MAPPING

        student = models.get_student_by_id_card(id_card)
        
        if not student:
            if name and _safe_str(row.get('学校类型', '')):
                result['valid'] = True
                result['status'] = 'new_record'
                result['name'] = name
                for chinese_field, db_field in field_mapping.items():
                    new_value = _safe_str(row.get(chinese_field, ''))
                    if not new_value:
                        continue
                    if db_field == '_attendance':
                        continue
                    result['changes'].append({
                        'field': chinese_field,
                        'db_field': db_field,
                        'old_value': '',
                        'new_value': new_value
                    })
            else:
                result['valid'] = False
                result['status'] = 'not_found'
                result['name'] = '-'
            results.append(result)
            continue
        
        result['name'] = student.get('name', '-')
        result['student_id'] = student.get('id')
        
        has_changes = False
        
        for chinese_field, db_field in field_mapping.items():
            new_value = _safe_str(row.get(chinese_field, ''))
            
            if not new_value:
                continue
            
            if db_field == '_attendance':
                id_card = student.get('id_card', '')
                att_status = today_attendance.get(id_card) if id_card else None
                if att_status == '出勤':
                    old_value = '出勤'
                elif att_status == '缺勤':
                    old_value = '缺勤'
                else:
                    old_value = ''
            elif db_field in ('school', 'enrollment_year'):
                old_value = str(student.get(db_field, '') or '')
                try:
                    if old_value and new_value:
                        if int(float(old_value)) == int(float(new_value)):
                            continue
                except (ValueError, TypeError):
                    pass
            else:
                old_value = str(student.get(db_field, ''))
            
            if new_value != old_value:
                result['changes'].append({
                    'field': chinese_field,
                    'db_field': db_field,
                    'old_value': old_value,
                    'new_value': new_value
                })
                has_changes = True
        
        if has_changes:
            result['status'] = 'pending'
        else:
            result['status'] = 'no_change'
        
        results.append(result)
    
    return results

def export_students_to_excel(students, fields=None):
    """导出学生toExcel。 参数: students, fields"""
    wb = Workbook()
    ws = wb.active
    ws.title = '学生信息'
    
    if fields is None:
        fields = config.ALL_STUDENT_FIELDS
    
    for col, field in enumerate(fields, 1):
        cell = ws.cell(row=1, column=col, value=field)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color='CCCCCC', end_color='CCCCCC', fill_type='solid')
        cell.alignment = Alignment(horizontal='center')
    
    for row_idx, student in enumerate(students, 2):
        for col_idx, field in enumerate(fields, 1):
            db_field = config.FIELD_MAPPING.get(field, field)
            
            if field == '年龄':
                value = calculate_age(student.get('birth_date'))
            elif field == '入学年份':
                v = student.get(db_field, 0)
                value = '' if v == 0 else v
            else:
                value = student.get(db_field, '')
            
            ws.cell(row=row_idx, column=col_idx, value=value)
    
    return wb

def cleanup_old_files(directory, max_age_hours=24):
    """清理指定目录下超过 max_age_hours 的旧文件。 参数: directory, max_age_hours"""
    now = datetime.now()
    for filename in os.listdir(directory):
        filepath = os.path.join(directory, filename)
        if os.path.isfile(filepath):
            file_time = datetime.fromtimestamp(os.path.getmtime(filepath))
            age = now - file_time
            if age.total_seconds() > max_age_hours * 3600:
                try:
                    os.remove(filepath)
                except OSError:
                    pass

def format_datetime(dt):
    """格式化datetime。 参数: dt"""
    if isinstance(dt, datetime):
        return dt.strftime('%Y-%m-%d %H:%M:%S')
    return str(dt)

def get_school_type_name(school_type):
    """获取学校type姓名。 参数: school_type"""
    return config.SCHOOL_TYPES.get(int(school_type), '未知')

def parse_filters_from_request(request):
    """解析筛选条件fromrequest。 参数: request"""
    filters = {}
    
    if request.args.get('school'):
        filters['school'] = int(request.args.get('school'))
    
    grade = request.args.get('grade')
    if grade:
        filters['grade'] = grade
    
    major = request.args.get('major')
    if major:
        filters['major'] = major
    
    admin_class = request.args.get('admin_class')
    if admin_class:
        filters['admin_class'] = admin_class
    
    virtual_class = request.args.get('virtual_class')
    if virtual_class:
        filters['virtual_class'] = virtual_class
    
    if request.args.get('registration_status'):
        filters['registration_status'] = request.args.get('registration_status')
    
    if request.args.get('attendance_status'):
        filters['attendance_status'] = request.args.get('attendance_status')
    
    if request.args.get('gender'):
        filters['gender'] = request.args.get('gender')

    if request.args.get('ethnicity'):
        filters['ethnicity'] = request.args.get('ethnicity')

    if request.args.get('search'):
        filters['search'] = request.args.get('search')

    student_no = request.args.get('student_no')
    if student_no:
        filters['student_no'] = student_no

    name = request.args.get('name')
    if name:
        filters['name'] = name
    
    min_age = request.args.get('min_age')
    if min_age:
        try:
            min_val = int(min_age)
            if 0 <= min_val <= 120:
                filters['min_age'] = str(min_val)
        except (ValueError, TypeError):
            pass

    max_age = request.args.get('max_age')
    if max_age:
        try:
            max_val = int(max_age)
            if 0 <= max_val <= 120:
                filters['max_age'] = str(max_val)
        except (ValueError, TypeError):
            pass

    has_photo = request.args.get('has_photo')
    if has_photo in ('0', '1'):
        filters['has_photo'] = has_photo

    major_direction = request.args.get('major_direction')
    if major_direction:
        filters['major_direction'] = major_direction

    household_type = request.args.get('household_type')
    if household_type:
        filters['household_type'] = household_type

    admission_type = request.args.get('admission_type')
    if admission_type:
        filters['admission_type'] = admission_type

    change_info = request.args.get('change_info')
    if change_info:
        filters['change_info'] = change_info

    created_from = (request.args.get('created_from') or '').strip()
    if created_from:
        try:
            datetime.strptime(created_from, '%Y-%m-%d')
            filters['created_from'] = created_from
        except (ValueError, TypeError):
            pass

    created_to = (request.args.get('created_to') or '').strip()
    if created_to:
        try:
            datetime.strptime(created_to, '%Y-%m-%d')
            filters['created_to'] = created_to
        except (ValueError, TypeError):
            pass

    return filters


def is_private_ip(ip):
    """判断 IP 是否为局域网私有地址。 参数: ip"""
    try:
        return ipaddress.IPv4Address(ip).is_private
    except Exception:
        return False


_cached_lan_ip = None


def get_lan_ip():
    """获取局域网 IP，结果缓存避免重复计算"""
    global _cached_lan_ip
    if _cached_lan_ip is not None:
        return _cached_lan_ip
    
    ip = _detect_lan_ip()
    _cached_lan_ip = ip
    return ip


def _detect_lan_ip():
    """探测本机局域网 IP（多种策略降级）。"""
    if platform.system() != 'Windows':
        try:
            s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            s.connect(('223.5.5.5', 53))
            ip = s.getsockname()[0]
            s.close()
            return ip
        except Exception:
            return '127.0.0.1'

    try:
        host_name = socket.gethostname()
        for addrinfo in socket.getaddrinfo(host_name, None, socket.AF_INET):
            ip = addrinfo[4][0]
            if is_private_ip(ip) and not ip.startswith('127.'):
                return ip
    except Exception as e:
        logger.debug("通过主机名获取局域网IP失败: %s", e, exc_info=True)

    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(('223.5.5.5', 53))
        ip = s.getsockname()[0]
        s.close()
        if is_private_ip(ip):
            return ip
    except Exception as e:
        logger.debug("通过UDP探测获取局域网IP失败，降级返回127.0.0.1: %s", e, exc_info=True)

    return '127.0.0.1'

def school_visible_contains(school_visible, school):
    """检查 school_visible 配置是否包含指定学校。 参数: school_visible, school"""
    if not school_visible or school_visible == '*':
        return True
    if not school:
        return True
    allowed = [s.strip() for s in school_visible.split(',') if s.strip()]
    return str(school) in allowed

_school_major_cache = {}

def _derive_school(major):
    """从已有学生数据推导学校类别。返回 (school, derived)"""
    major = (major or '').strip()
    if not major:
        return 1, False
    if major in _school_major_cache:
        return _school_major_cache[major], True
    try:
        conn = models.get_db()
        row = conn.execute(
            'SELECT DISTINCT school FROM students WHERE major=? AND school IS NOT NULL AND school>0 ORDER BY school LIMIT 1',
            (major,)
        ).fetchone()
    except Exception:
        row = None
    if row:
        val = row[0]
        _school_major_cache[major] = val
        return val, True
    return 1, False

def _get_card_reader_ws_url():
    """从配置文件读取读卡器 WebSocket 地址，不存在则返回默认值"""
    import json
    config_file = config.CARD_READER_CONFIG_FILE
    try:
        with open(config_file, 'r', encoding='utf-8') as f:
            data = json.load(f)
            return data.get('url', config.DEFAULT_CARD_READER_WS_URL)
    except (FileNotFoundError, json.JSONDecodeError):
        return config.DEFAULT_CARD_READER_WS_URL


def _set_card_reader_ws_url(url):
    """将读卡器 WebSocket 地址写入配置文件"""
    import json
    os.makedirs(os.path.dirname(config.CARD_READER_CONFIG_FILE), exist_ok=True)
    with open(config.CARD_READER_CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump({'url': url}, f, ensure_ascii=False)


def normalize_date_string(value):
    """将斜杠日期格式（YYYY/M/D、M/D/YY）统一转为 YYYY-MM-DD。

    已是 YYYY-MM-DD 格式的字符串原样返回；空值返回空字符串；
    无法识别的格式返回原值。
    """
    if not value or not isinstance(value, str):
        return value if value else ''
    value = value.strip()
    if not value:
        return ''
    # 已是标准格式 YYYY-MM-DD
    if re.match(r'^\d{4}-\d{2}-\d{2}$', value):
        return value
    # YYYY/M/D
    m = re.match(r'^(\d{4})/(\d{1,2})/(\d{1,2})$', value)
    if m:
        return f'{m.group(1)}-{m.group(2).zfill(2)}-{m.group(3).zfill(2)}'
    # M/D/YY
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{2})$', value)
    if m:
        return f'20{m.group(3)}-{m.group(1).zfill(2)}-{m.group(2).zfill(2)}'
    return value
