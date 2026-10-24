# -*- coding: utf-8 -*-
"""档案派遣：人才机构字典、区划匹配、批次生成与状态流转。

设计要点：
1. 一县一机构（talent_agency.county_code 唯一）；
2. 匹配按 县 → 市（先 [:4]+99，再 [:4]+00）→ 省（239900 / [:2]+0000）逐级回退；
3. 派遣记录存机构与打印信息快照，历史单据不随字典/配置变化；
4. 按县分批，批次号自动生成：{届}-{县名}-{两位序号}。
"""
import logging
import re
from datetime import datetime

import config
from models.db import get_db

logger = logging.getLogger(__name__)

STATUS_PENDING = '待派遣'
STATUS_SENT = '已派遣'
STATUS_RETURNED = '已退回'
STATUS_INVALID = '已失效'
STATUS_STUDY = '升学'

MATCH_COUNTY = '县级直寄'
MATCH_CITY = '市级代收'
MATCH_PROVINCE = '省级代收'
MATCH_MANUAL = '人工指定'
MATCH_NONE = '未匹配'

LEVEL_LABELS = {1: '省', 2: '市', 3: '县区'}
EXCEL_LEVEL_MAP = {'省': '省', '市、地区': '市', '市': '市', '县（区）': '县区', '县区': '县区'}


# ---------------------------------------------------------------- 基础工具
def norm_code(value):
    """区划代码归一化为 6 位：去非数字、12 位取前 6、4/2 位补齐。"""
    if value is None:
        return ''
    s = re.sub(r'\D', '', str(value))
    if not s:
        return ''
    if len(s) > 6:
        s = s[:6]
    if len(s) == 4:
        s += '00'
    elif len(s) == 2:
        s += '0000'
    return s.zfill(6)


def graduate_year_of(grade):
    """届 = 年级（4 位年份）+ 3；无法解析时退回系统学年。"""
    m = re.search(r'(\d{4})', str(grade or ''))
    if m:
        return str(int(m.group(1)) + 3)
    return str(getattr(config, 'CURRENT_SCHOOL_YEAR', datetime.now().year))


def _group_key(student):
    """学生归县键：县码优先；市辖区/县这类虚级改用市。返回 (code, name)。"""
    county_code = norm_code(student.get('district_county_code'))
    county_name = (student.get('district_county_name') or '').strip()
    city_code = norm_code(student.get('district_city_code'))
    city_name = (student.get('district_city_name') or '').strip()

    if county_code and county_name not in ('市辖区', '县', '省直辖县级行政区划'):
        return county_code, county_name or city_name
    if city_code:
        return city_code, city_name or county_name
    return '', county_name or ''


def _origin_place(student):
    """生源地：县名，缺县取市，缺市取省。"""
    for key in ('district_county_name', 'district_city_name', 'district_province_name'):
        val = (student.get(key) or '').strip()
        if val and val not in ('市辖区', '县'):
            return val
    return (student.get('district_province_name') or '').strip()


def _candidates(student):
    """按优先级生成 (代码, 层级) 候选：县 → 市 → 省。"""
    county_code = norm_code(student.get('district_county_code'))
    city_code = norm_code(student.get('district_city_code'))
    province_code = norm_code(student.get('district_province_code'))

    cands = []
    if county_code:
        cands.append((county_code, '县'))
    base = (county_code or city_code)[:4]
    if base:
        cands.append((base + '99', '市'))
        cands.append((base + '00', '市'))
    if city_code and all(city_code != c for c, _ in cands):
        cands.append((city_code, '市'))
    prov = (province_code or county_code or city_code)[:2]
    if prov:
        cands.append((prov + '9900', '省'))
        cands.append((prov + '0000', '省'))

    seen = set()
    result = []
    for code, level in cands:
        if code and code not in seen:
            seen.add(code)
            result.append((code, level))
    return result


def _match_agency(row, agency_map):
    """在已加载的机构字典里做三级回退，返回匹配信息字典。"""
    if not row.get('district_county_code') and not row.get('district_city_code'):
        return {'level': '', 'status': MATCH_NONE, 'reason': '区划缺失',
                'agency': None}
    if not norm_code(row.get('district_county_code')):
        fallback_reason = '县码缺失'
    else:
        fallback_reason = ''
    for code, level in _candidates(row):
        agency = agency_map.get(code)
        if not agency:
            continue
        if level == '县':
            status, reason = MATCH_COUNTY, ''
        elif level == '市':
            status, reason = MATCH_CITY, fallback_reason or '县无机构'
        else:
            status, reason = MATCH_PROVINCE, '市县均无机构'
        return {'level': level, 'status': status, 'reason': reason, 'agency': agency}
    return {'level': '', 'status': MATCH_NONE, 'reason': fallback_reason or '三级均无机构',
            'agency': None}


def _load_agency_map(conn):
    rows = conn.execute(
        'SELECT * FROM talent_agency WHERE status = 1 AND matchable = 1'
    ).fetchall()
    return {r['county_code']: dict(r) for r in rows}


# ---------------------------------------------------------------- 机构字典
def list_agencies(keyword='', level='', status=''):
    """机构字典列表（含停用）。"""
    conn = get_db()
    where = ['1 = 1']
    params = []
    if keyword:
        where.append('(agency_name LIKE ? OR county_name LIKE ? OR county_code LIKE ? OR agency_no LIKE ?)')
        like = f'%{keyword}%'
        params.extend([like, like, like, like])
    if level:
        where.append('level = ?')
        params.append(level)
    if status != '':
        where.append('status = ?')
        params.append(int(status))
    rows = conn.execute(
        f"SELECT * FROM talent_agency WHERE {' AND '.join(where)} "
        'ORDER BY level DESC, county_code'
    , params).fetchall()
    conn.close()
    return [dict(r) for r in rows]


def toggle_agency(agency_id):
    conn = get_db()
    row = conn.execute('SELECT status FROM talent_agency WHERE id = ?', (agency_id,)).fetchone()
    if not row:
        conn.close()
        return None
    new_status = 0 if row['status'] else 1
    conn.execute('UPDATE talent_agency SET status = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?',
                 (new_status, agency_id))
    conn.commit()
    conn.close()
    return new_status


def create_agency(data):
    """单条新增机构（手工补录）。返回新增记录，校验失败抛 ValueError。"""
    code = norm_code(data.get('county_code'))
    name = (data.get('agency_name') or '').strip()
    if len(code) != 6:
        raise ValueError('请填写 6 位行政区划代码')
    if not name:
        raise ValueError('请填写机构全称')
    agency_no = (data.get('agency_no') or '').strip() or code

    level = (data.get('level') or '').strip()
    district = _district_index()
    info = district.get(code)
    if info:
        county_name, admin_level = info
        level = level or LEVEL_LABELS.get(admin_level, '县区')
    else:
        level = level or _infer_level(code, '')
        if level == '市':
            county_name = district.get(code[:4] + '00', ('', 0))[0]
        elif level == '省':
            county_name = district.get(code[:2] + '0000', ('', 0))[0]
        else:
            county_name = ''
    if level not in ('省', '市', '县区'):
        level = '县区'

    conn = get_db()
    try:
        if conn.execute('SELECT 1 FROM talent_agency WHERE county_code = ?', (code,)).fetchone():
            raise ValueError(f'区划代码 {code} 已有机构，请停用原机构后再新增')
        if conn.execute('SELECT 1 FROM talent_agency WHERE agency_no = ?', (agency_no,)).fetchone():
            raise ValueError(f'机构编号 {agency_no} 已存在')
        cursor = conn.execute(
            'INSERT INTO talent_agency (agency_no, county_code, county_name, level, '
            'agency_name, address, phone, remark, matchable) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 1)',
            (agency_no, code, county_name, level, name,
             (data.get('address') or '').strip(), (data.get('phone') or '').strip(),
             (data.get('remark') or '').strip())
        )
        conn.commit()
        row = conn.execute('SELECT * FROM talent_agency WHERE id = ?', (cursor.lastrowid,)).fetchone()
        return dict(row)
    finally:
        conn.close()


def _district_index():
    """读取区划字典（6 位前缀 -> (name, level)）；失败返回空字典，导入不中断。"""
    try:
        from models.district import ensure_attached, DistrictDbMissing, import_dictionary
        conn = get_db()
        try:
            ensure_attached(conn)
        except DistrictDbMissing:
            import_dictionary()
            ensure_attached(conn)
        rows = conn.execute(
            'SELECT district_code, district_name, admin_level FROM dm.dm_district_code '
            'WHERE admin_level <= 3'
        ).fetchall()
        return {r['district_code'][:6]: (r['district_name'], r['admin_level']) for r in rows}
    except Exception as exc:  # noqa: BLE001
        logger.warning('区划字典不可用，机构导入跳过层级校验: %s', exc)
        return {}


def _history_map(conn):
    rows = conn.execute(
        'SELECT old_code, new_code FROM dm_history_mapping WHERE is_deleted = 0'
    ).fetchall()
    return {norm_code(r['old_code']): norm_code(r['new_code']) for r in rows if r['old_code']}


def _infer_level(code, raw_level=''):
    """区划字典查无此码时按编码规则推断层级：9900→省，99/00 结尾→市，其余→县区。

    机构表对「市」「省」使用自定义码（如 230199、239900），字典中并无这些记录，
    但它们正是匹配回退时要用的码，必须参与匹配。
    """
    if code.endswith('9900'):
        return '省'
    if code.endswith('99') or code.endswith('00'):
        return '市'
    return EXCEL_LEVEL_MAP.get(raw_level, '县区')


def import_agencies(file_obj):
    """导入人才机构 Excel（表头含「行政区划代码」）。

    返回统计：新增/更新/冲突/特殊区域/别名转换/未识别行。
    """
    from openpyxl import load_workbook

    wb = load_workbook(file_obj, data_only=True, read_only=True)
    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_idx = None
    headers = []
    for i, row in enumerate(raw_rows[:10]):
        vals = [str(c).strip() if c is not None else '' for c in row]
        if '行政区划代码' in vals:
            header_idx = i
            headers = vals
            break
    if header_idx is None:
        raise ValueError('未找到「行政区划代码」表头，请使用标准模板')

    def col(name):
        try:
            return headers.index(name)
        except ValueError:
            return None

    c_code = col('行政区划代码')
    c_no = col('机构编号')
    c_name = col('机构全称')
    c_level = col('机构层级')
    c_addr = col('通讯地址')
    c_phone = col('联系电话')
    if c_code is None or c_name is None:
        raise ValueError('模板至少需要「行政区划代码」「机构全称」两列')

    district = _district_index()
    conn = get_db()
    history = _history_map(conn)

    stats = {'inserted': 0, 'updated': 0, 'conflict': 0, 'special': 0,
             'aliased': 0, 'skipped': 0, 'conflicts': [], 'specials': []}

    def cell(row, idx):
        if idx is None or idx >= len(row):
            return ''
        val = row[idx]
        return str(val).strip() if val is not None else ''

    for row in raw_rows[header_idx + 1:]:
        raw_code = cell(row, c_code)
        name = cell(row, c_name)
        if not raw_code and not name:
            continue
        code = norm_code(raw_code)
        if not code or not name:
            stats['skipped'] += 1
            continue

        if code in history and history[code] and history[code] != code:
            code = history[code]
            stats['aliased'] += 1

        info = district.get(code)
        raw_level = cell(row, c_level)
        if info:
            county_name, admin_level = info
            level = LEVEL_LABELS.get(admin_level, '县区')
            matchable = 1 if admin_level <= 3 else 0
        else:
            level = _infer_level(code, raw_level)
            matchable = 1 if level in ('省', '市', '县区') else 0
            if level == '市':
                county_name = district.get(code[:4] + '00', ('', 0))[0]
            elif level == '省':
                county_name = district.get(code[:2] + '0000', ('', 0))[0]
            else:
                # 编码不属于任何行政区划 → 特殊区域，不参与自动匹配，仅支持人工指定
                county_name = ''
                matchable = 0
                stats['special'] += 1
                stats['specials'].append(f'{raw_code} {name}')

        agency_no = cell(row, c_no) or code
        address = cell(row, c_addr)
        phone = cell(row, c_phone)

        existing = conn.execute(
            'SELECT id, agency_no FROM talent_agency WHERE county_code = ?', (code,)
        ).fetchone()
        if existing and existing['agency_no'] != agency_no:
            stats['conflict'] += 1
            stats['conflicts'].append(f'{code} 已被机构「{existing["agency_no"]}」占用，跳过「{name}」')
            continue

        if existing:
            conn.execute(
                'UPDATE talent_agency SET agency_no = ?, county_name = ?, level = ?, '
                'agency_name = ?, address = ?, phone = ?, matchable = ?, status = 1, '
                'updated_at = CURRENT_TIMESTAMP WHERE id = ?',
                (agency_no, county_name, level, name, address, phone, matchable, existing['id'])
            )
            stats['updated'] += 1
        else:
            conn.execute(
                'INSERT INTO talent_agency (agency_no, county_code, county_name, level, '
                'agency_name, address, phone, matchable) VALUES (?, ?, ?, ?, ?, ?, ?, ?)',
                (agency_no, code, county_name, level, name, address, phone, matchable)
            )
            stats['inserted'] += 1

    conn.commit()
    conn.close()
    return stats


def agency_template_rows():
    """机构导入模板表头与示例行。"""
    return (
        ['行政区划代码', '机构编号', '机构全称', '机构层级', '通讯地址', '联系电话'],
        [['230102', '23010201', '哈尔滨市道里区人才服务中心', '县（区）',
          '哈尔滨市道里区XX街1号', '0451-00000000']]
    )


# ---------------------------------------------------------------- 转递单配置
def get_all_settings():
    conn = get_db()
    rows = conn.execute('SELECT * FROM archive_settings ORDER BY school').fetchall()
    conn.close()
    return {r['school']: dict(r) for r in rows}


def save_settings(school, data, operator=''):
    conn = get_db()
    conn.execute(
        'INSERT INTO archive_settings (school, school_name, handler, handler_phone, '
        'receipt_address, receipt_postal, updated_by, updated_at) '
        'VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP) '
        'ON CONFLICT(school) DO UPDATE SET school_name = excluded.school_name, '
        'handler = excluded.handler, handler_phone = excluded.handler_phone, '
        'receipt_address = excluded.receipt_address, receipt_postal = excluded.receipt_postal, '
        'updated_by = excluded.updated_by, updated_at = CURRENT_TIMESTAMP',
        (int(school), data.get('school_name', ''), data.get('handler', ''),
         data.get('handler_phone', ''), data.get('receipt_address', ''),
         data.get('receipt_postal', ''), operator)
    )
    conn.commit()
    conn.close()


# ---------------------------------------------------------------- 毕业生归县
def load_graduates(grade=None, school=None):
    """加载毕业生并附带归县键、机构匹配结果、当前派遣状态。"""
    conn = get_db()
    sql = '''
        SELECT s.id, s.student_no, s.id_card, s.name, s.school, s.grade, s.admin_class,
               s.student_phone, s.parent_phone, s.major_direction,
               s.district_province_code, s.district_province_name,
               s.district_city_code, s.district_city_name,
               s.district_county_code, s.district_county_name,
               d.id AS dispatch_id, d.dispatch_status, d.batch_no,
               d.agency_name, d.agency_address, d.agency_phone
        FROM students s
        LEFT JOIN archive_dispatch d ON d.student_id = s.id AND d.is_deleted = 0
        WHERE s.is_deleted = 0 AND s.registration_status = '毕业'
    '''
    params = []
    if grade:
        sql += ' AND s.grade = ?'
        params.append(grade)
    if school:
        sql += ' AND s.school = ?'
        params.append(int(school))
    sql += ' ORDER BY s.grade DESC, s.admin_class, s.name'
    rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
    agency_map = _load_agency_map(conn)
    conn.close()

    for row in rows:
        code, name = _group_key(row)
        row['group_code'] = code
        row['group_name'] = name or '未匹配区划'
        row['origin_place'] = _origin_place(row)
        match = _match_agency(row, agency_map)
        row['match_level'] = match['level']
        row['match_status'] = match['status']
        row['match_reason'] = match['reason']
        row['matched_agency'] = match['agency']
        row['dispatch_status'] = row.get('dispatch_status') or ''
    return rows


def summarize_counties(rows):
    """按县汇总：待派遣/已派遣/升学/未匹配/可派遣。"""
    groups = {}
    for row in rows:
        key = row['group_code']
        item = groups.setdefault(key, {
            'county_code': key,
            'county_name': row['group_name'],
            'total': 0, 'pending': 0, 'dispatchable': 0, 'sent': 0,
            'study': 0, 'unmatched': 0, 'invalid': 0, 'returned': 0,
            'match_status': row['match_status'],
            'agency_name': (row['matched_agency'] or {}).get('agency_name', ''),
        })
        item['total'] += 1
        status = row['dispatch_status']
        if status == STATUS_STUDY:
            item['study'] += 1
        elif status == STATUS_SENT:
            item['sent'] += 1
        elif status == STATUS_INVALID:
            item['invalid'] += 1
        elif status == STATUS_RETURNED:
            # 退档待重派：仍属可办范围，但单独计数，避免与「从未派过」混淆
            item['returned'] += 1
            if row['matched_agency']:
                item['dispatchable'] += 1
            else:
                item['unmatched'] += 1
        else:
            item['pending'] += 1
            if row['matched_agency']:
                item['dispatchable'] += 1
            else:
                item['unmatched'] += 1

    result = list(groups.values())
    result.sort(key=lambda x: (x['county_code'] == '', x['county_name']))
    return result


def coverage(grade=None):
    """机构字典覆盖率：毕业生涉及县数 / 已配机构县数 / 缺口清单。"""
    groups = summarize_counties(load_graduates(grade))
    gaps = [g for g in groups if g['county_code'] and not g['agency_name']]
    covered = [g for g in groups if g['county_code'] and g['agency_name']]
    return {
        'group_total': len([g for g in groups if g['county_code']]),
        'covered': len(covered),
        'gaps': gaps,
    }


# ---------------------------------------------------------------- 批次生成
def _next_batch_no(conn, graduate_year, county_name):
    prefix = f'{graduate_year}-{county_name}-'
    rows = conn.execute(
        'SELECT batch_no FROM archive_dispatch WHERE batch_no LIKE ?', (prefix + '%',)
    ).fetchall()
    max_seq = 0
    for r in rows:
        m = re.search(r'-(\d+)$', r['batch_no'] or '')
        if m:
            max_seq = max(max_seq, int(m.group(1)))
    return f'{prefix}{max_seq + 1:02d}'


def create_batch(group_code, student_ids, operator, overrides=None):
    """按县生成一批派遣记录（默认待派遣），返回批次号与统计。"""
    overrides = overrides or {}
    targets = set(int(x) for x in student_ids)
    rows = [r for r in load_graduates() if r['group_code'] == group_code and r['id'] in targets]
    if not rows:
        raise ValueError('所选学生不属于该县或已不是毕业生')
    # 一个批次只能属于一个校别：批次级操作（整批标记已派遣/导出）会作用于全批，
    # 混校别会导致另一所学校被连带处理，破坏「两校分别处理」的口径
    if len({r.get('school') for r in rows}) > 1:
        raise ValueError('所选学生包含多个校别，请先按校别筛选后分别生成批次')

    settings = get_all_settings()
    graduate_year = graduate_year_of(rows[0].get('grade'))
    county_name = rows[0]['group_name']

    conn = get_db()
    try:
        conn.execute('BEGIN IMMEDIATE')
        batch_no = _next_batch_no(conn, graduate_year, county_name)
        inserted = 0
        skipped = []
        for row in rows:
            if row['dispatch_status'] in (STATUS_SENT, STATUS_STUDY):
                skipped.append(f"{row['name']}：已派遣/升学，跳过")
                continue
            if not row['matched_agency']:
                skipped.append(f"{row['name']}：未匹配到机构，跳过")
                continue
            existing = conn.execute(
                'SELECT id FROM archive_dispatch WHERE student_id = ? AND is_deleted = 0',
                (row['id'],)
            ).fetchone()
            if existing:
                skipped.append(f"{row['name']}：已存在派遣记录，跳过")
                continue

            agency = row['matched_agency']
            cfg = settings.get(row.get('school') or 1, {})
            school_name = overrides.get('school_name') or cfg.get('school_name', '')
            handler = overrides.get('handler') or cfg.get('handler', '')
            handler_phone = overrides.get('handler_phone') or cfg.get('handler_phone', '')
            receipt_address = cfg.get('receipt_address', '')
            receipt_postal = cfg.get('receipt_postal', '')

            conn.execute(
                '''INSERT INTO archive_dispatch (
                    batch_no, dispatch_no, student_id, student_no, id_card, student_name,
                    graduate_year, school, grade, admin_class, county_code, county_name,
                    origin_place, student_phone, parent_phone,
                    match_level, match_status, match_reason,
                    agency_id, agency_no, agency_name, agency_address, agency_phone,
                    school_name, handler, handler_phone, receipt_address, receipt_postal,
                    dispatch_status, operator
                ) VALUES (?, '', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                (batch_no, row['id'], row.get('student_no') or '', row.get('id_card') or '',
                 row.get('name') or '', graduate_year, row.get('school'), row.get('grade') or '',
                 row.get('admin_class') or '', row['group_code'], row['group_name'],
                 row['origin_place'], row.get('student_phone') or '', row.get('parent_phone') or '',
                 row['match_level'], row['match_status'], row['match_reason'],
                 agency.get('id'), agency.get('agency_no') or '', agency.get('agency_name') or '',
                 agency.get('address') or '', agency.get('phone') or '',
                 school_name, handler, handler_phone, receipt_address, receipt_postal,
                 STATUS_PENDING, operator)
            )
            inserted += 1
        # 生成批次即预填转递日期（打印转递单时即需显示日期）；
        # 后续标记「已派遣」时会被实际派遣日期覆盖
        if inserted:
            conn.execute(
                "UPDATE archive_dispatch SET dispatch_date = ? "
                "WHERE batch_no = ? AND is_deleted = 0 "
                "AND (dispatch_date IS NULL OR dispatch_date = '')",
                (datetime.now().strftime('%Y-%m-%d'), batch_no)
            )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {'batch_no': batch_no, 'inserted': inserted, 'skipped': skipped,
            'county_name': county_name, 'graduate_year': graduate_year}


def list_batches(school=None):
    """批次汇总：退档、重派、升学、失效分别计数。"""
    conn = get_db()
    sql = """SELECT batch_no, county_code, county_name,
                  COUNT(*) AS total,
                  SUM(CASE WHEN dispatch_status = '已派遣' THEN 1 ELSE 0 END) AS sent,
                  SUM(CASE WHEN dispatch_status = '已退回' THEN 1 ELSE 0 END) AS returned,
                  SUM(CASE WHEN dispatch_status = '已派遣'
                            AND COALESCE(returned_at, '') != '' THEN 1 ELSE 0 END) AS redispatch,
                  SUM(CASE WHEN dispatch_status = '升学' THEN 1 ELSE 0 END) AS study,
                  SUM(CASE WHEN dispatch_status = '已失效' THEN 1 ELSE 0 END) AS invalid,
                  MIN(created_at) AS created_at,
                  MIN(handler) AS handler
           FROM archive_dispatch
           WHERE is_deleted = 0 AND batch_no != ''"""
    params = []
    if school:
        sql += ' AND school = ?'
        params.append(int(school))
    sql += ' GROUP BY batch_no ORDER BY created_at DESC'
    rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
    conn.close()
    for row in rows:
        # 应派遣 = 总人数 - 升学 - 已失效；否则升学/失效会一直占着分母，批次永远「未完成」
        row['active_total'] = (row['total'] or 0) - (row['study'] or 0) - (row['invalid'] or 0)
    return rows


def get_batch_rows(batch_no):
    conn = get_db()
    rows = conn.execute(
        'SELECT * FROM archive_dispatch WHERE batch_no = ? AND is_deleted = 0 ORDER BY id',
        (batch_no,)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


# ---------------------------------------------------------------- 状态流转
def update_dispatch(dispatch_id, action, payload=None, operator=''):
    """状态流转：sent / returned / reopen。"""
    payload = payload or {}
    conn = get_db()
    row = conn.execute('SELECT * FROM archive_dispatch WHERE id = ? AND is_deleted = 0',
                       (dispatch_id,)).fetchone()
    if not row:
        conn.close()
        raise ValueError('派遣记录不存在')

    status = row['dispatch_status']
    if action == 'sent':
        if status not in (STATUS_PENDING, STATUS_RETURNED):
            conn.close()
            raise ValueError(f'当前状态「{status}」不能标记为已派遣')
        date = (payload.get('dispatch_date') or datetime.now().strftime('%Y-%m-%d')).strip()
        conn.execute(
            'UPDATE archive_dispatch SET dispatch_status = ?, dispatch_date = ?, '
            'operator = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?',
            (STATUS_SENT, date, operator, dispatch_id)
        )
    elif action == 'returned':
        if status != STATUS_SENT:
            conn.close()
            raise ValueError('只有已派遣记录可以标记退回')
        # 记录退档时间：非空即「曾退档」，重派后仍保留，用于筛出重派单
        conn.execute(
            "UPDATE archive_dispatch SET dispatch_status = ?, dispatch_date = '', "
            'returned_at = ?, operator = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?',
            (STATUS_RETURNED, datetime.now().strftime('%Y-%m-%d'), operator, dispatch_id)
        )
    elif action == 'reopen':
        if status != STATUS_RETURNED:
            conn.close()
            raise ValueError('只有已退回记录可以恢复待派遣')
        conn.execute(
            'UPDATE archive_dispatch SET dispatch_status = ?, operator = ?, '
            'updated_at = CURRENT_TIMESTAMP WHERE id = ?',
            (STATUS_PENDING, operator, dispatch_id)
        )
    else:
        conn.close()
        raise ValueError('未知操作')
    conn.commit()
    conn.close()


def update_batch_status(batch_no, action, operator=''):
    """整批状态流转：整批标记已派遣（仅待派遣；已退回须先恢复待派遣再派）。

    返回 {count, skipped_returned}，让前端提示「退档记录未被一并改状态」。
    """
    if action != 'sent':
        raise ValueError('未知操作')
    conn = get_db()
    try:
        skipped = conn.execute(
            'SELECT COUNT(1) FROM archive_dispatch WHERE batch_no = ? AND is_deleted = 0 '
            'AND dispatch_status = ?', (batch_no, STATUS_RETURNED)
        ).fetchone()[0]
        cur = conn.execute(
            'UPDATE archive_dispatch SET dispatch_status = ?, dispatch_date = ?, '
            'operator = ?, updated_at = CURRENT_TIMESTAMP '
            'WHERE batch_no = ? AND is_deleted = 0 AND dispatch_status = ?',
            (STATUS_SENT, datetime.now().strftime('%Y-%m-%d'), operator, batch_no,
             STATUS_PENDING)
        )
        conn.commit()
        count = cur.rowcount
    finally:
        conn.close()
    if not count:
        raise ValueError('该批次没有可标记为已派遣的待派遣记录')
    return {'count': count, 'skipped_returned': skipped}


def update_dispatch_agency(dispatch_id, agency_id=None, manual=None):
    """人工指定接收机构（manual 为直接填写机构信息）。"""
    conn = get_db()
    row = conn.execute('SELECT * FROM archive_dispatch WHERE id = ? AND is_deleted = 0',
                       (dispatch_id,)).fetchone()
    if not row:
        conn.close()
        raise ValueError('派遣记录不存在')
    if row['dispatch_status'] not in (STATUS_PENDING, STATUS_RETURNED):
        conn.close()
        raise ValueError('只有待派遣/已退回记录可以改机构')

    if agency_id:
        agency = conn.execute('SELECT * FROM talent_agency WHERE id = ?', (agency_id,)).fetchone()
        if not agency:
            conn.close()
            raise ValueError('机构不存在')
        conn.execute(
            'UPDATE archive_dispatch SET agency_id = ?, agency_no = ?, agency_name = ?, '
            'agency_address = ?, agency_phone = ?, match_level = ?, match_status = ?, '
            'match_reason = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?',
            (agency['id'], agency['agency_no'], agency['agency_name'], agency['address'],
             agency['phone'], '人工', MATCH_MANUAL, '', dispatch_id)
        )
    else:
        manual = manual or {}
        conn.execute(
            'UPDATE archive_dispatch SET agency_id = NULL, agency_no = ?, agency_name = ?, '
            'agency_address = ?, agency_phone = ?, match_level = ?, match_status = ?, '
            'match_reason = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?',
            (manual.get('agency_no', ''), manual.get('agency_name', ''),
             manual.get('address', ''), manual.get('phone', ''),
             '人工', MATCH_MANUAL, '', dispatch_id)
        )
    conn.commit()
    conn.close()


def list_students_for_agency(keyword=''):
    conn = get_db()
    like = f'%{keyword}%'
    rows = conn.execute(
        'SELECT id, agency_no, county_code, county_name, level, agency_name, address, phone, matchable '
        'FROM talent_agency WHERE status = 1 AND (agency_name LIKE ? OR county_name LIKE ? OR county_code LIKE ?) '
        'ORDER BY county_code LIMIT 50', (like, like, like)
    ).fetchall()
    conn.close()
    return [dict(r) for r in rows]


# ---------------------------------------------------------------- 升学
def mark_further_study(student_ids, operator=''):
    """标记升学：无记录则新建；待派遣/已退回则转升学并清空机构与派遣日期。"""
    ids = [int(x) for x in student_ids]
    if not ids:
        return {'marked': 0, 'skipped': []}
    rows = {r['id']: r for r in load_graduates() if r['id'] in set(ids)}
    settings = get_all_settings()
    conn = get_db()
    marked, skipped = 0, []
    try:
        conn.execute('BEGIN IMMEDIATE')
        for sid in ids:
            row = rows.get(sid)
            if not row:
                skipped.append(f'ID={sid}：不是毕业生，跳过')
                continue
            current = conn.execute(
                'SELECT id, dispatch_status FROM archive_dispatch WHERE student_id = ? AND is_deleted = 0',
                (sid,)
            ).fetchone()
            if current:
                if current['dispatch_status'] == STATUS_STUDY:
                    skipped.append(f"{row['name']}：已是升学，跳过")
                    continue
                if current['dispatch_status'] == STATUS_SENT:
                    skipped.append(f"{row['name']}：已派遣，不能直接转升学")
                    continue
                conn.execute(
                    "UPDATE archive_dispatch SET dispatch_status = ?, agency_id = NULL, "
                    "agency_no = '', agency_name = '', agency_address = '', agency_phone = '', "
                    "dispatch_date = '', match_level = '', match_status = '无需派遣', "
                    "match_reason = '升学', operator = ?, updated_at = CURRENT_TIMESTAMP "
                    'WHERE id = ?',
                    (STATUS_STUDY, operator, current['id'])
                )
                marked += 1
                continue

            cfg = settings.get(row.get('school') or 1, {})
            conn.execute(
                '''INSERT INTO archive_dispatch (
                    batch_no, dispatch_no, student_id, student_no, id_card, student_name,
                    graduate_year, school, grade, admin_class, county_code, county_name,
                    origin_place, student_phone, parent_phone,
                    match_level, match_status, match_reason,
                    school_name, handler, handler_phone, receipt_address, receipt_postal,
                    dispatch_status, operator
                ) VALUES ('', '', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '', '无需派遣', '升学',
                          ?, ?, ?, ?, ?, ?, ?)''',
                (sid, row.get('student_no') or '', row.get('id_card') or '', row.get('name') or '',
                 graduate_year_of(row.get('grade')), row.get('school'), row.get('grade') or '',
                 row.get('admin_class') or '', row['group_code'], row['group_name'],
                 row['origin_place'], row.get('student_phone') or '', row.get('parent_phone') or '',
                 cfg.get('school_name', ''), cfg.get('handler', ''), cfg.get('handler_phone', ''),
                 cfg.get('receipt_address', ''), cfg.get('receipt_postal', ''),
                 STATUS_STUDY, operator)
            )
            marked += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return {'marked': marked, 'skipped': skipped}


def cancel_further_study(student_ids):
    ids = [int(x) for x in student_ids]
    if not ids:
        return 0
    conn = get_db()
    placeholders = ','.join('?' for _ in ids)
    cur = conn.execute(
        f"UPDATE archive_dispatch SET is_deleted = 1, updated_at = CURRENT_TIMESTAMP "
        f"WHERE dispatch_status = ? AND is_deleted = 0 AND student_id IN ({placeholders})",
        [STATUS_STUDY] + ids
    )
    conn.commit()
    count = cur.rowcount
    conn.close()
    return count


def list_further_study(grade=None, county='', school=None):
    conn = get_db()
    sql = '''SELECT d.*, s.major_direction FROM archive_dispatch d
             LEFT JOIN students s ON s.id = d.student_id
             WHERE d.is_deleted = 0 AND d.dispatch_status = ?'''
    params = [STATUS_STUDY]
    if grade:
        sql += ' AND d.grade = ?'
        params.append(grade)
    if county:
        sql += ' AND d.county_code = ?'
        params.append(county)
    if school:
        sql += ' AND d.school = ?'
        params.append(int(school))
    sql += ' ORDER BY d.grade DESC, d.county_code, d.student_name'
    rows = [dict(r) for r in conn.execute(sql, params).fetchall()]
    conn.close()
    return rows


def mark_dispatches_invalid(student_ids, operator=''):
    """毕业撤销联动：把相关有效派遣记录标为已失效（留痕不删）。"""
    ids = [int(x) for x in student_ids]
    if not ids:
        return 0
    conn = get_db()
    placeholders = ','.join('?' for _ in ids)
    cur = conn.execute(
        f"UPDATE archive_dispatch SET dispatch_status = ?, operator = ?, "
        f"updated_at = CURRENT_TIMESTAMP WHERE is_deleted = 0 AND dispatch_status != ? "
        f"AND student_id IN ({placeholders})",
        [STATUS_INVALID, operator, STATUS_INVALID] + ids
    )
    conn.commit()
    count = cur.rowcount
    conn.close()
    return count


# ---------------------------------------------------------------- 导出
# 不出转递单的状态：已退回（档案已退回，待改机构重派）、已失效（毕业已撤销）、
# 升学（档案转录取院校，无需派遣）
EXPORT_EXCLUDED_STATUSES = (STATUS_RETURNED, STATUS_INVALID, STATUS_STUDY)


def build_export_rows(batch_no, redispatch_only=False):
    """组装 docx 模板填充所需的行数据（含年月日拆分）。

    redispatch_only=True：只取「曾退档且已重新派遣」的记录，用于打印重派单；
    默认：取待派遣 + 已派遣，排除退档/失效/升学。
    """
    rows = get_batch_rows(batch_no)
    result = []
    for row in rows:
        status = row.get('dispatch_status')
        if redispatch_only:
            if status != STATUS_SENT or not (row.get('returned_at') or '').strip():
                continue
        elif status in EXPORT_EXCLUDED_STATUSES:
            continue
        date = (row.get('dispatch_date') or '').strip()
        # 历史批次兼容：旧数据生成时未写 dispatch_date（created_at 为 UTC，
        # 直接取会差一天）；未派遣/已退回的按导出当天（本地）填充
        if not date and row.get('dispatch_status') in (STATUS_PENDING, STATUS_RETURNED):
            date = datetime.now().strftime('%Y-%m-%d')
        year = month = day = ''
        m = re.match(r'(\d{4})-(\d{1,2})-(\d{1,2})', date)
        if m:
            year, month, day = str(int(m.group(1))), str(int(m.group(2))), str(int(m.group(3)))
        item = dict(row)
        item.update({
            'dispatch_year': year,
            'dispatch_month': month,
            'dispatch_day': day,
        })
        result.append(item)
    return result
